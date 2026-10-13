#!/usr/bin/env python3
"""Master ticker 2026-06-30 ~06:53 — TT short-link FIXED, move to completed, update summary."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# === TASK 1: Mark TT short-link task as COMPLETE and move to Completed section ===
open_idx = None
completed_idx = None
tt_para = None
empty_bullet_exists = False

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text == "Open":
        open_idx = i
    elif text == "Completed":
        completed_idx = i
    if open_idx and not completed_idx:
        if "TT short-link" in text or "/TT/" in text:
            tt_para = p
        if text == "[]":
            empty_bullet_exists = True

if tt_para:
    for run in tt_para.runs:
        run.text = ""
    tt_para.runs[0].text = (
        "[x] FIXED — TT short-link restored. The /TT/ shortcut now correctly "
        "redirects all sub-paths (including play links) to their /TchoTcho/ equivalents. "
        "The issue was that only the bare /TT and /TT/ paths had redirect rules; "
        "sub-paths like /TT/play/[game] were falling through to a 404. "
        "Added a catch-all redirect rule on the server."
    )

    elem = tt_para._element
    parent = elem.getparent()
    parent.remove(elem)

    for p in d.paragraphs:
        if p.text.strip() == "Completed":
            p._element.addnext(elem)
            break
    print("OK: TT short-link moved to Completed")
else:
    print("WARNING: TT short-link task not found in Open")

# Ensure empty [] bullet still exists in Open
if not empty_bullet_exists:
    for p in d.paragraphs:
        if p.text.strip() == "Open":
            new_p = OxmlElement('w:p')
            pPr = p._element.find(qn('w:pPr'))
            # Find the next paragraph with a bullet to copy its formatting
            next_elem = p._element.getnext()
            if next_elem is not None:
                next_pPr = next_elem.find(qn('w:pPr'))
                if next_pPr is not None:
                    new_p.insert(0, deepcopy(next_pPr))
            run_elem = OxmlElement('w:r')
            text_elem = OxmlElement('w:t')
            text_elem.text = "[]"
            run_elem.append(text_elem)
            new_p.append(run_elem)
            # Insert after Open heading but find the right spot (after In Progress heading)
            # Actually we just need one [] in Open section
            p._element.addnext(new_p)
            print("OK: added empty [] bullet in Open")
            break

# === TASK 2: Update bulleted summary ===
in_summary = False
summary_bullets = []
summary_heading_para = None

for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text and not in_summary:
        in_summary = True
        summary_heading_para = p
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
        elif text in ("Open", "In Progress", "Completed") or text.startswith("["):
            break

# Find "Old summary bullets" section
old_summary_para = None
for p in d.paragraphs:
    if "Old summary bullets" in p.text:
        old_summary_para = p
        break

# If we have 3+ summary bullets, move the oldest to old-summaries
if len(summary_bullets) >= 3 and old_summary_para is not None:
    oldest_para = summary_bullets[-1][1]
    oldest_text = oldest_para.text.strip()

    new_p = OxmlElement('w:p')
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_para._element.addnext(new_p)

    oldest_para._element.getparent().remove(oldest_para._element)
    print(f"OK: moved oldest summary to old-summaries")

# Re-scan to find insert point for new summary bullet
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

NEW_SUMMARY = (
    "2026-06-30 ~06:53 CDT (master ticker): TT short-link FIXED — "
    "all /TT/ paths now correctly redirect to /TchoTcho/ (play links included). "
    "Both in-progress tasks remain blocked on user input. Admin queue empty."
)

if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = NEW_SUMMARY
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)
    print("OK: new summary bullet inserted")

d.save(DOC)
print("SAVED")
