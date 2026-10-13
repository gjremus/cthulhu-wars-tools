#!/usr/bin/env python3
"""Master ticker 2026-06-30 ~09:11b — cleanup: fix Open section to have empty [], remove empty paras."""
import docx
from docx.oxml.ns import qn

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)
body = d.element.body

# Find key indices
open_idx = None
inprog_idx = None

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text == "Open":
        open_idx = i
    elif text == "In Progress":
        inprog_idx = i

print(f"Open={open_idx}, InProg={inprog_idx}")

# Fix paragraph 40: should be just "[]"
p40 = d.paragraphs[open_idx + 1]
print(f"Para 40 before: {repr(p40.text[:100])}")
for run in p40.runs:
    run.text = ""
p40.runs[0].text = "[]"
print(f"Para 40 after: '[]'")

# Remove empty paragraphs between Open and In Progress (indices 41-51 are all empty)
paras_to_remove = []
for i in range(open_idx + 2, inprog_idx):
    p = d.paragraphs[i]
    if not p.text.strip():
        paras_to_remove.append(p._element)

print(f"Removing {len(paras_to_remove)} empty paragraphs from Open section")
for elem in paras_to_remove:
    body.remove(elem)

# Also remove empty paragraph in summary (index 38)
for i, p in enumerate(d.paragraphs):
    if p.text.strip() == "Bulleted summary (newest at top)":
        # Check paragraphs between summary and Open for empties
        for j in range(i + 1, 100):
            pj = d.paragraphs[j]
            if pj.text.strip() == "Open":
                break
            if not pj.text.strip():
                body.remove(pj._element)
                print(f"Removed empty summary paragraph")
                break
        break

d.save(DOC)
print("SAVED")
