#!/usr/bin/env python3
# Library ticker 2026-06-24 ~22:03 — routine healthy pass, nothing changed/deployed.
#  1. Trash oldest backup (handled before this script runs).
#  2. Create new backup alongside docx.
#  3. Insert one brief grovelling health bullet at the top of the summary.
#  4. Demote the prior ~13:33 nothing-actionable health bullet to the Old section
#     (prefix OLD:) so the displayed summary stays within MAX-3 and MAX-2-nothing-actionable.
import docx, shutil, time, os

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


# ── 1. Insert one brief health summary bullet at the top of the summary ───────
NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~22:03. Grovelling thanks — re-read the full rules, drained the game "
    "console’s message box (empty, nothing new), and confirmed the live site healthy: the Library "
    "single-player edition and the More-Neutral-Units beta both answered instantly. I walked every open "
    "and waiting job; nothing was changed or deployed, everything substantial still waits only on your "
    "go-aheads, and the saved copies of this sheet stay at the allowed three."
)
ANCHOR = "Last tick: 2026-06-24 ~13:33"
anchor_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR):
        anchor_p = p
        break
assert anchor_p is not None, "13:33 summary anchor missing"
anchor_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_p.style)

# ── 2. Demote the prior ~13:33 health bullet to the Old section ───────────────
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break
assert old_heading is not None, "Old summary heading missing"

if not anchor_p.text.strip().startswith("OLD:"):
    set_text(anchor_p, "OLD: " + anchor_p.text.strip())
old_heading._p.addnext(anchor_p._p)

d.save(DOC)
print("OK: health bullet added ; 13:33 bullet demoted ; backup=%s" % os.path.basename(bak))
