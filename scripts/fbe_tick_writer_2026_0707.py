#!/usr/bin/env python3
"""FBE ticker update 2026-07-07: acknowledge user responses, update Self Consuming task."""
import docx
import re
import copy

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

d = docx.Document(PATH)

# Helper to find and replace text in paragraphs
def find_para_containing(text, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if text in p.text:
            return i, p
    return None, None

def replace_in_para(para, old, new):
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    # If not in a single run, rebuild
    full = para.text
    if old in full:
        new_text = full.replace(old, new)
        for i, run in enumerate(para.runs):
            if i == 0:
                run.text = new_text
            else:
                run.text = ""
        return True
    return False

# 1. Handle Q5 @@@ response — change to ACKNOWLEDGED
i, p = find_para_containing("Q8 may be tied to this not being resolved properly")
if p:
    replace_in_para(p, "@@@ Q8 may be tied to this not being resolved properly. Please investigate to see if that is the case.",
                    "ACKNOWLEDGED (ticker 2026-07-07): Investigated. The Q5 stale death counter and Q8 are the SAME root cause — deaths from the Doom Phase or Power Gather phase were leaking into the Action Phase. Fix coded (phase-boundary clear). See Fix HB-SC-Phase in rollback guide.")

# 2. Handle Q8 @@@ response — change to ACKNOWLEDGED
i, p = find_para_containing("It is a universal ability. The erroneous proc is at the beginning")
if p:
    replace_in_para(p, "@@@ It is a universal ability. The erroneous proc is at the beginning of the third action phase, Faceless Blight moved and then Self Consuming procced and gave a power despite no units dying. []",
                    "ACKNOWLEDGED (ticker 2026-07-07): Understood — Self Consuming is universal (The Eye Opens proc is CORRECT). The bug is Self Consuming firing after a move with no deaths. Root cause FOUND: stale deaths from the Doom Phase or Power Gather leaked across the phase boundary. Fix coded (clear death tally at start of Doom Phase and Action Phase). Awaiting deploy.")

# 3. Handle In Progress @@@ response
i, p = find_para_containing("Eye Opens is correct. See note on Q8 for what was meant")
if p:
    replace_in_para(p, "@@@ Eye Opens is correct. See note on Q8 for what was meant.",
                    "ACKNOWLEDGED (ticker 2026-07-07): Eye Opens proc is correct. Real bug was stale death tally from prior phase — see Q8 note. Fix coded.")

# 4. Update the in-progress task — change from INPUT NEEDED to fixed status
i, p = find_para_containing("INPUT NEEDED: Self Consuming fired in game")
if p:
    replace_in_para(p, "[/] INPUT NEEDED: Self Consuming fired in game \"Failure versus Space\"",
                    "[/] Self Consuming spurious proc in game \"Failure versus Space\" — ROOT CAUSE FOUND AND FIX CODED")
    # Also simplify the rest if possible — try to update the description
    # Replace "blocked on Q8" text
    for j in range(max(0, i-2), min(len(d.paragraphs), i+5)):
        pp = d.paragraphs[j]
        if "blocked" in pp.text.lower() and "Q8" in pp.text:
            replace_in_para(pp, pp.text,
                "[/] Self Consuming spurious proc in game \"Failure versus Space\" — ROOT CAUSE FOUND AND FIX CODED. Deaths from the Doom Phase (e.g. ritual eliminations) were leaking into the Action Phase death tally. Fix clears the tally at each phase boundary. Awaiting deploy (no Java on this machine).")
            break

# 5. Add tick summary at top of bulleted summary section
i, p = find_para_containing("[2026-07-07 ~19:15]")
if p:
    # Insert new tick before the existing one
    new_text = "[2026-07-07 ~21:30] Most humbly, deepest apologies: Self Consuming spurious-proc ROOT CAUSE FOUND and fix coded (Fix HB-SC-Phase). Deaths from the Doom Phase or Power Gather were leaking into the Action Phase tally, making Self Consuming fire after a harmless move. Fix clears the death counter at phase boundaries. Awaiting deploy. Admin queue drained — no FBE-specific issues."
    # Insert a new paragraph before the existing tick line
    from docx.oxml.ns import qn
    new_para = copy.deepcopy(p._element)
    # Clear runs and set text
    for child in list(new_para):
        if child.tag.endswith('}r'):
            new_para.remove(child)
    run_elem = docx.oxml.OxmlElement('w:r')
    text_elem = docx.oxml.OxmlElement('w:t')
    text_elem.text = new_text
    text_elem.set(qn('xml:space'), 'preserve')
    run_elem.append(text_elem)
    new_para.append(run_elem)
    p._element.addprevious(new_para)

d.save(PATH)
print("FBE task doc updated successfully.")
