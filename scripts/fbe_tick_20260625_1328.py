"""
FBE tick 2026-06-25 ~13:28 — update doc:
- Update the "nothing happened" summary bullet at para 36 with current timestamp and corrected facts
  (worktree has an uncommitted Ghatanothoa change, not FBE — no FBE work needed this tick)
- Leave Open section with one empty [] as required
"""
import sys
from docx import Document
from docx.oxml.ns import qn
import copy

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks.docx"

doc = Document(DOC_PATH)

# Find the "nothing happened" / quiet tick bullet — it's para 36 (0-indexed: paragraph index 36)
# Text: "[2026-06-25 ~13:12] Humbly reporting another quiet tick..."
# We update it in place.

OLD_STUB = "[2026-06-25 ~13:12]"
NEW_TEXT = ("[2026-06-25 ~13:28] Humbly reporting another quiet tick with apologies for the earlier inaccuracy: "
            "admin queue is empty, Homebrew site healthy (cache 20260624-125953, HTTP 200), no open Faceless Blight tasks remain. "
            "Note: the worktree holds an uncommitted correction to a different faction's overlay text — not Faceless Blight, so no Faceless Blight action is needed.")

updated = False
for para in doc.paragraphs:
    if OLD_STUB in para.text:
        # Clear all runs and set new text in first run
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = NEW_TEXT
        else:
            para.add_run(NEW_TEXT)
        updated = True
        print(f"Updated para: {para.text[:80]}")
        break

if not updated:
    print("ERROR: could not find the target paragraph to update")
    sys.exit(1)

doc.save(DOC_PATH)
print("SAVED OK")
