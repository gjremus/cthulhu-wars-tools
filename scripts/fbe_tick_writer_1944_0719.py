#!/usr/bin/env python3
"""FBE ticker 2026-07-19 ~19:44 — update nothing-happened summary timestamp."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx")

d = docx.Document(DOC)

NEW_TEXT = ("[2026-07-19 ~19:44] Most humbly, deepest apologies: "
            "admin queue inaccessible (token file locked by Google Drive). "
            "No open or in-progress FBE tasks remain. "
            "Your worthless servant grovels abjectly.")

# Update the top-of-doc tick line
top_updated = False
for p in d.paragraphs:
    if p.text.strip().startswith("[2026-07-19 ~18:00]") and "admin queue" in p.text:
        for run in p.runs:
            run.text = ""
        p.runs[0].text = NEW_TEXT
        top_updated = True
        print("OK: updated top-of-doc tick line")
        break

# Update the bulleted summary (first matching entry)
in_summary = False
summary_updated = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary = True
        continue
    if in_summary and p.text.strip().startswith("[2026-07-19") and "admin queue" in p.text:
        for run in p.runs:
            run.text = ""
        p.runs[0].text = NEW_TEXT
        summary_updated = True
        print("OK: updated bulleted summary tick line")
        break

if not top_updated and not summary_updated:
    print("ERROR: could not find any matching bullet to update")
    exit(1)

d.save(DOC)
print("SAVED")
