#!/usr/bin/env python3
"""Fix the two READ: paragraphs that had curly-apostrophe search misses."""

import docx
import sys
from lxml import etree
from docx.oxml.ns import qn

DOCX_PATH = ("/Users/gremus/Library/CloudStorage/"
             "GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/"
             "Cthulhu Wars/HomeBrews/Faceless Blight Current Tasks.docx")


def set_para_text(para, new_text):
    for r in para._element.findall(qn('w:r')):
        para._element.remove(r)
    r_elem = etree.SubElement(para._element, qn('w:r'))
    t_elem = etree.SubElement(r_elem, qn('w:t'))
    t_elem.text = new_text
    t_elem.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def main():
    doc = docx.Document(DOCX_PATH)

    fixed = 0
    for i, para in enumerate(doc.paragraphs):
        t = para.text
        if 'READ:' in t and 'Anything' in t and 'Color' in t:
            print(f"Para {i}: {repr(t[:80])}")
            set_para_text(para,
                "ADDRESSED (ticker 2026-06-25): Reference game reviewed; "
                "root causes identified and fixed (Self Consuming scope, "
                "Overlord of Death action type, Succor phase, Animated Rush "
                "carry logic). All confirmed live.")
            fixed += 1
        elif 'READ:' in t and 'Succor' in t and 'Doom Phase' in t:
            print(f"Para {i}: {repr(t[:80])}")
            set_para_text(para,
                "ADDRESSED (ticker 2026-06-25): Succor fixed — fires only "
                "in the Doom Phase, returns to the Doom menu after resolving "
                "(does not consume your main action), and has a "
                "once-per-Doom-Phase guard. Fixed and deployed.")
            fixed += 1

    print(f"Fixed {fixed} paragraphs.")
    doc.save(DOCX_PATH)
    print("Saved.")

    # Verify no READ: markers remain in Q section
    doc2 = docx.Document(DOCX_PATH)
    remaining = [(i, p.text[:100]) for i, p in enumerate(doc2.paragraphs)
                 if p.text.startswith('READ:')]
    if remaining:
        print(f"WARNING: {len(remaining)} READ: markers still present:")
        for idx, txt in remaining:
            print(f"  para {idx}: {txt}")
    else:
        print("Confirmed: no remaining READ: markers in the document.")


if __name__ == "__main__":
    main()
