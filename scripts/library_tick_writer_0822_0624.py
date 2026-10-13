#!/usr/bin/env python3
# Library ticker 2026-06-24 ~08:22 — the live site has RECOVERED.
#  1. Mark the 05:18 outage in-progress item [/] -> [x] RESOLVED (keep original for history).
#  2. Add one brief recovery summary bullet at top.
#  3. Demote the three outage summary bullets (05:46/05:18/05:03) to the Old section
#     so the displayed summary stays within the MAX-3 cap.
import docx, shutil, time, os

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


# ── 1. Mark the outage in-progress item resolved ──────────────────────────────
OUT_PREFIX = "[/] INPUT NEEDED — the live multiplayer site is DOWN AGAIN"
RESOLVED_PREFIX = (
    "[x] RESOLVED 2026-06-24 ~08:22 — grovelling apologies for the long alarm; the live site is back up "
    "and answering normally (the Library single-player edition, the More-Neutral-Units beta and a known-good "
    "game page all responded instantly this pass, and the admin message box is empty). The outage that began "
    "around 05:18 has fully cleared and needs nothing further from you; I changed and deployed nothing. "
    "Original alarm kept for history: "
)
outage_done = False
for p in d.paragraphs:
    if p.text.strip().startswith(OUT_PREFIX):
        set_text(p, RESOLVED_PREFIX + p.text.strip())
        outage_done = True
        break
assert outage_done, "outage in-progress anchor missing"

# ── 2. Insert one brief recovery summary bullet at top of summary ─────────────
NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~08:22. Grovelling relief and apologies — the live site is BACK UP: the Library "
    "single-player edition, the More-Neutral-Units beta and a known-good game page all answered instantly, "
    "and the admin message box is empty. The earlier multi-hour outage has fully cleared and needs nothing "
    "further from you; I changed and deployed nothing, and everything substantial still waits only on your go-aheads."
)
B546 = "Last tick: 2026-06-24 ~05:46"
b546 = None
for p in d.paragraphs:
    if p.text.strip().startswith(B546):
        b546 = p
        break
assert b546 is not None, "05:46 summary anchor missing"
b546.insert_paragraph_before(NEW_SUMMARY, style=b546.style)

# ── 3. Demote the three outage summary bullets to the Old section ─────────────
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break
assert old_heading is not None, "Old summary heading missing"

# Process oldest->newest so the newest ends up directly under the heading.
DEMOTE = [
    "Last tick: 2026-06-24 ~05:03",
    "Last tick: 2026-06-24 ~05:18",
    "Last tick: 2026-06-24 ~05:46",
]
for prefix in DEMOTE:
    target = None
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            target = p
            break
    if target is None:
        continue
    if not target.text.strip().startswith("OLD:"):
        set_text(target, "OLD: " + target.text.strip())
    # Physically move the paragraph element to just after the Old heading.
    old_heading._p.addnext(target._p)

d.save(DOC)
print("OK: outage resolved=%s ; recovery bullet added ; 3 outage bullets demoted ; backup=%s"
      % (outage_done, os.path.basename(bak)))
