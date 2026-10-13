#!/usr/bin/env python3
"""Backup ticker 2026-06-29 ~08:17 CDT — MF1/Prime Cause bug root-caused and fixed."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# ── 1. Move the MF1/Prime Cause task from Open [] to In Progress [/] ─────────
# Find the Open task about "Million Favored Ones" and change its [] to [/]
# Also move it to after the "In Progress" heading

open_mf1_idx = None
open_mf1_para = None
in_progress_idx = None
open_empty_bullet_idx = None

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if "Million Favored Ones" in text and text.startswith("[]"):
        open_mf1_idx = i
        open_mf1_para = p
    if text == "In Progress":
        in_progress_idx = i
    # Track the empty [] bullet (leave one in Open)
    if text == "[]" and open_mf1_idx is not None and open_empty_bullet_idx is None:
        open_empty_bullet_idx = i

if open_mf1_para is not None:
    # Change [] to [/] and update text with root cause info
    old_text = open_mf1_para.text
    new_text = old_text.replace("[] ", "[/] ", 1).replace(
        "Investigating.",
        "ROOT CAUSE FOUND (backup ticker): after Prime Cause replaces a unit, "
        "the battle's internal tracking list still contains the old unit and lacks "
        "the new one, so Million Favored Ones sees stale data. Fix applied to all "
        "4 builds (BB, TT, MNU, HB). Cannot compile on this machine (no Java) — "
        "awaiting compilation and deploy."
    )
    # Clear existing runs and set new text
    for run in open_mf1_para.runs:
        run.text = ""
    if open_mf1_para.runs:
        open_mf1_para.runs[0].text = new_text
    else:
        open_mf1_para.text = new_text

    # Move it: remove from Open, insert after "In Progress" heading
    if in_progress_idx is not None:
        in_progress_para = d.paragraphs[in_progress_idx]
        # Clone the paragraph element and insert after In Progress heading
        elem = open_mf1_para._element
        parent = elem.getparent()
        cloned = deepcopy(elem)
        in_progress_para._element.addnext(cloned)
        # Remove the original from Open
        parent.remove(elem)

# ── 2. Add summary bullet ────────────────────────────────────────────────────
in_summary = False
first_summary = None
summary_bullets = []

for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text and not in_summary:
        in_summary = True
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
            if first_summary is None:
                first_summary = p
        elif text in ("Open", "In Progress", "Completed") or text.startswith("["):
            break

# Move oldest (3rd+) summary bullet to old summaries
old_summary_idx = None
for i, p in enumerate(d.paragraphs):
    if "Old summary bullets" in p.text:
        old_summary_idx = i
        break

if len(summary_bullets) >= 3 and old_summary_idx is not None:
    oldest_para = summary_bullets[2][1]
    oldest_text = oldest_para.text.strip()
    old_summary_para = d.paragraphs[old_summary_idx]
    new_p = OxmlElement('w:p')
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_para._element.addnext(new_p)
    oldest_para._element.getparent().remove(oldest_para._element)

# Re-find first summary after potential removal
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

# Insert new summary bullet before the first existing one
if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = (
        "2026-06-29 ~08:17 CDT (backup ticker): MF1/Prime Cause bug ROOT-CAUSED "
        "and FIXED in all 4 builds. The battle tracker wasn't updated when Prime "
        "Cause swapped a unit, so post-battle abilities (Million Favored Ones) saw "
        "stale data. Cannot compile (no Java on this machine) — code fix ready, "
        "awaiting build and deploy."
    )
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)

d.save(DOC)
print("OK: MF1/Prime Cause task moved to In Progress, summary bullet added")
