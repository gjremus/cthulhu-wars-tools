#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~16:08 — safety-net pass, nothing changed/deployed.
#  1. Insert one brief grovelling health bullet at the top of the summary.
#  2. Demote the prior ~15:52 nothing-actionable bullet to the Old section
#     (prefix OLD:) so the displayed summary stays within MAX-3 and MAX-2-nothing-actionable.
#  3. Rotate oldest backup file to trash if needed (max 3 copies).
import docx, shutil, time, os, glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

# Rotate: keep only 3 backups, move oldest to trash
bak_files = sorted(glob.glob(DOC + ".bak.pre_write_*"))
while len(bak_files) > 3:
    oldest = bak_files.pop(0)
    trash = os.path.expanduser("~/.Trash/")
    shutil.move(oldest, os.path.join(trash, os.path.basename(oldest)))
    print("Trashed old backup: %s" % os.path.basename(oldest))

d = docx.Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~16:08 (backup ticker). Grovelling apologies — re-read the full rules, "
    "drained the admin message box (empty, nothing new), and confirmed all five games fully healthy "
    "(Library, More-Neutral-Units, TchoTcho, Bubastis, and Homebrew all answered instantly). Walked "
    "every open and waiting job; nothing was changed or deployed, and everything still waits only on "
    "your go-aheads (chiefly the sideways-map correction for the More-Neutral-Units edition you "
    "actually see, and the crashed five-player game fix)."
)

# The most recent bullet to demote
DEMOTE_CANDIDATES = [
    "Last tick: 2026-06-25 ~15:52 (backup ticker)",
    "Last tick: 2026-06-25 ~15:39 (backup ticker)",
    "Last tick: 2026-06-25 ~14:48 (backup ticker)",
]

anchor_p = None
for candidate in DEMOTE_CANDIDATES:
    for p in d.paragraphs:
        if p.text.strip().startswith(candidate):
            anchor_p = p
            break
    if anchor_p is not None:
        break

assert anchor_p is not None, "Could not find recent summary bullet to demote"

# Insert new bullet before the one we're demoting
anchor_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_p.style)

# Demote the found bullet to Old section
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break
assert old_heading is not None, "Old summary heading missing"

if not anchor_p.text.strip().startswith("OLD:"):
    set_text(anchor_p, "OLD: " + anchor_p.text.strip())
old_heading._p.addnext(anchor_p._p)

d.save(DOC)
print("OK: health bullet added; prior bullet demoted; backup=%s" % os.path.basename(bak))
