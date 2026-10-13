#!/usr/bin/env python3
"""Append the HB Fix 116 entry to the Homebrew Deployed Online Rollback Guide.docx,
matching the existing Heading1 / Heading2 + Body Text structure of prior fixes.
Append-only (add_paragraph); never regex-deletes."""
import docx

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/HomeBrews/Homebrew Deployed Online Rollback Guide.docx"

d = docx.Document(PATH)
before = len(d.paragraphs)

def h1(t): d.add_paragraph(t, style="Heading 1")
def h2(t): d.add_paragraph(t, style="Heading 2")
def body(t): d.add_paragraph(t, style="Body Text")

h1("Fix 116 (2026-06-14) — Custodian / Librarian Agony die roll now shows in the log on replay and online games")

h2("Symptom.")
body("When the Librarian or Custodian rolled the special Agony die, the roll-result line (“Agony Die rolled N”, or “… + 1 for staying = N total”) appeared in live local play but was MISSING from online (multiplayer/host) games and from any reloaded save — those only showed the later “assigned N Agony to <faction>” line, so a player reviewing the log could not see what was actually rolled.")

h2("Root cause.")
body("The roll-result log was written INSIDE the RollAgony live callback (the `x => { log(...); <AssignAction> }` lambda in solo/MapExpansion.scala). Game replay short-circuits RollAgony: on a recorded roll it parses the recorded resolved action and applies it directly (solo/CthulhuWarsSolo.scala ~line 600-604), so the live callback — and therefore the log() inside it — never runs during replay/online reconstruction. Only the assign action it produced was recorded, and the assign action did not re-emit the roll line. This is the same class of bug already solved for combat rolls via the “Ghroth” pattern (the roll log lives inside a RECORDED action's perform so it replays).")

h2("Fix.")
body("Moved the roll-result log out of the live callback and into two new RECORDED ForcedActions, mirroring GhrothRollAction. (1) Added case classes CustodianAgonyRolledAction(self, r, rolled, staying) and LibrarianAgonyRolledAction(self, r, rolled, staying) alongside the existing Library-expansion action case classes in solo/MapExpansion.scala. (2) The four RollAgony callbacks (Custodian move/stay, Librarian move/stay) no longer log inline; each now returns the matching ...AgonyRolledAction carrying the rolled value and the staying flag. (3) Added two handlers in LibraryExpansion.perform that emit the verbatim existing log wording — non-staying: “Agony Die rolled N”; staying: “Agony Die rolled N + 1 for staying = N+1 total” — then Force the existing CustodianAssignAgonyAction / LibrarianAssignAgonyAction with exactly the same arguments they received before (Custodian remaining = rolled or rolled+1; Librarian total=remaining = rolled or rolled+1). Because the new rolled-action is what gets recorded and replayed (reflection-based serializer, no registration needed — same mechanism as GhrothRollAction), the roll line now replays for online games and reloaded saves. No change to CthulhuWarsSolo.scala was required: its existing RollAgony replay path parses and performs whatever action the callback returns, which is now the rolled-action.")

h2("Files changed.")
body("solo/MapExpansion.scala ONLY — two new ForcedAction case classes (CustodianAgonyRolledAction, LibrarianAgonyRolledAction); four RollAgony callbacks changed to return the new rolled-action instead of logging inline; two new perform handlers that log the roll (verbatim prior wording) then Force the unchanged assign action. No serializer edit needed (reflection-based: lookupClass(“cws.” + name) in solo/Serialize.scala). solo/CthulhuWarsSolo.scala unchanged.")
body("NOT included in this build (pre-existing, in-progress, does NOT compile — same as Fix 113/114/115): the in-flight Faceless Blight (FBE) faction WIP across solo/Battle.scala, solo/BotFBE.scala, solo/FactionFBE.scala, solo/Game.scala. These were git-stashed (“TEMP park FBE WIP for Agony-log build”) before building so the shipped HB build contains ONLY the Fix 116 change on top of Fix 115, then restored to the working tree afterward (owner WIP and the pre-existing FBE stash untouched).")
body("Version bumped: solo/common.sbt Homebrew v 2.6 → Homebrew v 2.7.")

h2("Build + ship.")
body("Stashed the 4 FBE-WIP files, then: cd /Users/gremus/Claude-Projects/cw-homebrew-wt/solo && sbt fullOptJS (Closure: 0 error(s), 0 warning(s)), then ./server-deploy/deploy-homebrew-to-vm.sh. Cache tag 20260614-193909, public HEAD HTTP 200 at https://cwo.freeddns.org/HB/. Restored the FBE WIP to the working tree after deploy.")

h2("Git commit.")
body("Not committed by the build agent (deploy was the authorized action; commit/push left to the owner per workflow). Working tree on branch homebrew carries the Fix 116 change in solo/MapExpansion.scala + the version bump in solo/common.sbt, alongside the untouched 4-file FBE WIP. The live HB bundle reflects Fix 116.")

h2("Rollback.")
body("Revert solo/MapExpansion.scala to its pre-Fix-116 state (restore the inline-log RollAgony callbacks; remove the two CustodianAgonyRolledAction / LibrarianAgonyRolledAction case classes and their two perform handlers) and solo/common.sbt to Homebrew v 2.6, rebuild Homebrew (sbt fullOptJS, with FBE WIP stashed), redeploy. Restores Fix 115 behavior — the Agony roll line shows in live local play but is absent from replay/online games.")

d.save(PATH)
after = len(d.paragraphs)
print("Appended Fix 116 entry. Paragraphs:", before, "->", after, "(+%d)" % (after - before))
