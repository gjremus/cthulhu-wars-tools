#!/usr/bin/env python3
"""Insert the 'nothing happened' bullet after Summary History header."""
import docx
from docx.oxml.ns import qn
import sys
import copy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

BULLET_TEXT = (
    "2026-06-27 ~19:53 (master ticker): No new tasks or admin entries. "
    "Admin queue empty. Server healthy (HTTP 200). Both in-progress items "
    "(Alternative Ancients, remote access) still awaiting input."
)

d = docx.Document(DOC)

header = None
header_idx = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip().startswith("Summary History"):
        header = p
        header_idx = i
        break

if header is None:
    print("ERROR: could not find Summary History header", file=sys.stderr)
    sys.exit(1)

# Find the next paragraph after header to copy its style
next_p = d.paragraphs[header_idx + 1] if header_idx + 1 < len(d.paragraphs) else None

# Create a new paragraph element after the header
new_p_elem = copy.deepcopy(next_p._element) if next_p else docx.oxml.OxmlElement('w:p')
# Clear the runs
for run in new_p_elem.findall(qn('w:r')):
    new_p_elem.remove(run)

# Add a run with our text
new_run = docx.oxml.OxmlElement('w:r')
# Copy run properties from the next paragraph's first run if available
if next_p and next_p.runs:
    rPr = next_p.runs[0]._element.find(qn('w:rPr'))
    if rPr is not None:
        new_run.append(copy.deepcopy(rPr))
new_text = docx.oxml.OxmlElement('w:t')
new_text.text = BULLET_TEXT
new_text.set(qn('xml:space'), 'preserve')
new_run.append(new_text)
new_p_elem.append(new_run)

# Insert after header
header._element.addnext(new_p_elem)

d.save(DOC)
print("OK: inserted 'nothing happened' bullet after Summary History header")
