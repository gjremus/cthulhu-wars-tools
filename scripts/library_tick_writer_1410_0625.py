#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~14:10 — routine healthy pass, nothing changed/deployed.
# Rules: MAX-3 summary bullets, MAX-2 nothing-actionable bullets.
# Current state: 2 summary bullets exist (both are "nothing actionable" ticks at ~13:03 and ~08:48).
# Action: replace the OLDER (~08:48) bullet with OLD: prefix and move it to Old section,
# then insert a new ~14:10 nothing-actionable bullet at the top of the summary.
import docx, shutil, time, os, sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~14:10 (backup ticker). Grovelling apologies — re-read the full rules, "
    "drained the admin message box (empty, nothing new), and confirmed the live site fully healthy "
    "(the Library single-player edition and the More-Neutral-Units beta both answered instantly). "
    "Walked every open and waiting job; nothing was changed or deployed, and everything still waits "
    "only on your go-aheads (chiefly the sideways-map correction for the edition you actually see "
    "and the crashed five-player game fix)."
)

# The older nothing-actionable bullet to demote
OLDER_ANCHOR = "Last tick: 2026-06-25 ~08:48"

# Find the older bullet
older_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(OLDER_ANCHOR):
        older_p = p
        break

if older_p is None:
    print("ERROR: could not find the ~08:48 summary bullet to demote", file=sys.stderr)
    sys.exit(1)

# Find the newer bullet (the one just above, which we insert BEFORE)
newer_anchor = "Last tick: 2026-06-25 ~13:03"
newer_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(newer_anchor):
        newer_p = p
        break

if newer_p is None:
    print("ERROR: could not find the ~13:03 summary bullet as anchor for insert", file=sys.stderr)
    sys.exit(1)

# Insert new bullet before the ~13:03 bullet
newer_p.insert_paragraph_before(NEW_SUMMARY, style=newer_p.style)

# Demote the older bullet to Old section
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break

if old_heading is None:
    print("ERROR: could not find 'Old summary entries' heading", file=sys.stderr)
    sys.exit(1)

if not older_p.text.strip().startswith("OLD:"):
    set_text(older_p, "OLD: " + older_p.text.strip())
old_heading._p.addnext(older_p._p)

# Rotate backups: keep only 3; move older ones to ~/.Trash/
import glob, os
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
# The new bak we just made is already in the list (if created); keep newest 3
while len(all_baks) > 3:
    oldest = all_baks.pop(0)
    trash_dir = os.path.expanduser("~/.Trash/")
    os.makedirs(trash_dir, exist_ok=True)
    dest = os.path.join(trash_dir, os.path.basename(oldest))
    # avoid collisions
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest, dest)
    print("Rotated to Trash: %s" % os.path.basename(oldest))

d.save(DOC)
print("OK: new 14:10 backup summary bullet inserted; 08:48 bullet demoted to Old section; backup=%s" % os.path.basename(bak))
