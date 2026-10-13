#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~21:39 — routine healthy pass, nothing changed.
# Updates the top summary bullet with a fresh timestamp.
import docx, shutil, time, os

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~21:39 (backup ticker). Grovelling apologies — re-read the full rules, "
    "drained the admin message box (empty, nothing new), and confirmed all five games fully healthy "
    "(Library, More-Neutral-Units, TchoTcho, Bubastis, and Homebrew all answered). Walked every "
    "open and waiting job; nothing was changed or deployed, and everything still waits only on your "
    "go-aheads (chiefly the sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

# Find summary bullets
summary_bullets = []
summary_start_idx = None
for i, p in enumerate(d.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Bulleted summary"):
        summary_start_idx = i
        continue
    if summary_start_idx is not None and txt.startswith(("Last tick:", "OLD:")):
        summary_bullets.append((i, p))
    if summary_start_idx is not None and txt.startswith("Open"):
        break

assert len(summary_bullets) >= 1, f"Expected >=1 summary bullets, found {len(summary_bullets)}"

# Replace the first (newest) bullet with our new text
first_bullet = summary_bullets[0][1]
set_text(first_bullet, NEW_SUMMARY)

d.save(DOC)
print("OK: library backup ticker health bullet written; backup=%s" % os.path.basename(bak))

# Enforce max 3 bak copies
doc_dir = os.path.dirname(DOC)
baks = sorted([f for f in os.listdir(doc_dir) if f.startswith("library build current tasks.docx.bak.")])
while len(baks) > 3:
    oldest = baks.pop(0)
    oldest_path = os.path.join(doc_dir, oldest)
    trash = os.path.expanduser("~/.Trash/")
    shutil.move(oldest_path, os.path.join(trash, oldest))
    print(f"Moved old backup to Trash: {oldest}")
