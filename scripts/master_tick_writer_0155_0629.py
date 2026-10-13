#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~01:55 — update 'nothing happened' summary bullet."""
import docx
import sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

NEW_TEXT = (
    "2026-06-29 ~01:55 CDT (master ticker): No new tasks or admin entries. "
    "Both in-progress items still blocked on creator input "
    "(Alternative Ancients clarifications, remote access preference)."
)

CANDIDATE_PREFIXES = [
    "2026-06-29 ~01:32 CDT (master ticker): No new tasks",
    "2026-06-29 ~01:24 CDT (backup ticker): No new tasks",
]

d = docx.Document(DOC)
target = None
for prefix in CANDIDATE_PREFIXES:
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            target = p
            break
    if target is not None:
        break

if target is None:
    print("ERROR: could not find 'nothing happened' bullet to update", file=sys.stderr)
    sys.exit(1)

if target.runs:
    target.runs[0].text = NEW_TEXT
    for extra in target.runs[1:]:
        extra.text = ""
else:
    target.add_run(NEW_TEXT)

d.save(DOC)
print(f"OK: updated master 'nothing happened' bullet to ~01:55 CDT")
