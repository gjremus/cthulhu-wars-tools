#!/usr/bin/env python3
"""
library_tick_writer_1303_0625.py
Library ticker pass, 2026-06-25 ~13:03.

Actions this pass:
1. Create .bak backup of docx.
2. Insert new summary bullet at top of Bulleted Summary.
3. Per Rule 21: demote the oldest nothing-actionable summary bullet (08:37) to OLD section.
4. Regenerate the .last_seen_tasks.md mirror.

No task status changes — all INPUT NEEDED items still waiting on user go-ahead.
Admin queue: empty. Site health: all four pages HTTP 200.
"""
import sys
import shutil
import time
import os
from docx import Document

LIB = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno"
)
DOC = os.path.join(LIB, "library build current tasks.docx")
MIRROR = os.path.join(LIB, ".last_seen_tasks.md")
TRASH_BAK = os.path.join(LIB, ".Trash-bak-rotation")

# ── Pre-flight: 2-minute mtime lockout ───────────────────────────────────────
file_mtime = os.path.getmtime(DOC)
age = time.time() - file_mtime
if age < 120:
    print(f"LOCKED: doc modified {age:.0f}s ago (lockout 120s)", file=sys.stderr)
    sys.exit(1)
print(f"Lockout clear: doc is {age:.0f}s old")

# ── Step 1: Rotate backups (keep max 3; oldest → .Trash-bak-rotation) ────────
bak_prefix = "library build current tasks.docx.bak.pre_write_"
existing_baks = sorted([
    f for f in os.listdir(LIB)
    if f.startswith(bak_prefix)
])
print(f"Existing backups: {len(existing_baks)}")
while len(existing_baks) >= 3:
    oldest = os.path.join(LIB, existing_baks.pop(0))
    dest = os.path.join(TRASH_BAK, os.path.basename(oldest))
    shutil.move(oldest, dest)
    print(f"Trashed oldest backup: {os.path.basename(oldest)}")

bak_path = os.path.join(LIB, bak_prefix + str(int(time.time())))
shutil.copy2(DOC, bak_path)
print(f"Backup created: {os.path.basename(bak_path)}")

# ── Step 2: Open the document ─────────────────────────────────────────────────
d = Document(DOC)

NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~13:03. Grovelling apologies — re-read the full rules, "
    "drained the admin message box (empty, nothing new), and confirmed the live site "
    "fully healthy (the Library single-player edition, the More-Neutral-Units beta, "
    "the admin console and a known-good game page all answered instantly). Walked every "
    "open and waiting job; nothing was changed or deployed, and everything still waits "
    "only on your go-aheads (chiefly the sideways-map correction for the More-Neutral-"
    "Units edition you actually see, and the crashed five-player game fix)."
)

ANCHOR_0848 = "Last tick: 2026-06-25 ~08:48"  # current top bullet
ANCHOR_0837 = "Last tick: 2026-06-25 ~08:37"  # oldest nothing-actionable to demote

# ── Step 3: Find anchor paragraphs ───────────────────────────────────────────
para_0848 = None
para_0837 = None
old_heading_para = None

for p in d.paragraphs:
    txt = p.text.strip()
    if txt.startswith(ANCHOR_0848):
        para_0848 = p
    if txt.startswith(ANCHOR_0837):
        para_0837 = p
    if txt == "Old summary entries (tick history)":
        old_heading_para = p

if para_0848 is None:
    print("ERROR: 08:48 anchor not found", file=sys.stderr)
    sys.exit(1)
if para_0837 is None:
    print("ERROR: 08:37 anchor not found", file=sys.stderr)
    sys.exit(1)
if old_heading_para is None:
    print("ERROR: 'Old summary entries' heading not found", file=sys.stderr)
    sys.exit(1)

print("All anchors found.")

# ── Step 4: Insert new summary before 08:48 bullet ───────────────────────────
def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)

para_0848.insert_paragraph_before(NEW_SUMMARY, style=para_0848.style)
print("Inserted new summary bullet")

# ── Step 5: Demote 08:37 bullet to OLD section ───────────────────────────────
# Re-find 08:37 paragraph after insertion shifted indices
para_0837_found = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR_0837):
        para_0837_found = p
        break

if para_0837_found is None:
    print("ERROR: 08:37 paragraph lost after insertion", file=sys.stderr)
    sys.exit(1)

if not para_0837_found.text.strip().startswith("OLD:"):
    set_text(para_0837_found, "OLD: " + para_0837_found.text.strip())

# Move to just after the old_heading_para
old_heading_para._p.addnext(para_0837_found._p)
print("Demoted 08:37 bullet to Old section")

# ── Step 6: Save ─────────────────────────────────────────────────────────────
d.save(DOC)
print(f"Saved: {DOC}")

# ── Step 7: Regenerate mirror ────────────────────────────────────────────────
lines = []
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lines.append("## " + p.text)
    else:
        lines.append(p.text)
with open(MIRROR, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Mirror regenerated: {len(lines)} paragraphs")
print("DONE")
