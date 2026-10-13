#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~15:52 — routine healthy pass, nothing changed/deployed.
# Current docx state (after 15:39 backup tick): visible bullets = [15:39, 15:22, 13:03].
# Adding 15:52 makes 4; demote 13:03 to OLD to stay at MAX-3.
import docx
import shutil
import time
import os
import sys
import glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~15:52 (backup ticker). Grovelling apologies — re-read "
    "the full rules, drained the admin message box (empty, nothing new), and confirmed "
    "all five games fully healthy (Library, More-Neutral-Units, TchoTcho, Bubastis, and "
    "Homebrew all answered instantly). Walked every open and waiting job; nothing was "
    "changed or deployed, and everything still waits only on your go-aheads (chiefly "
    "the sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

ANCHOR_15_39 = "Last tick: 2026-06-25 ~15:39"
DEMOTE_ANCHORS = [
    "Last tick: 2026-06-25 ~13:03",
    "Last tick: 2026-06-25 ~08:48",
]

# Find the insert anchor and the bullet to demote
anchor_p = None
demote_p = None

for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(ANCHOR_15_39) and not t.startswith("OLD:"):
        anchor_p = p
    for da in DEMOTE_ANCHORS:
        if t.startswith(da) and not t.startswith("OLD:"):
            if demote_p is None:
                demote_p = p
            break

if anchor_p is None:
    # Fallback: try 15:22 as anchor
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Last tick: 2026-06-25 ~15:22") and not t.startswith("OLD:"):
            anchor_p = p
            break

if anchor_p is None:
    print("ERROR: could not find any insert anchor paragraph", file=sys.stderr)
    sys.exit(1)

# Insert new summary bullet before the anchor
anchor_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_p.style)

# Demote the oldest visible bullet to OLD:
if demote_p is not None:
    full_text = demote_p.text
    # Clear all runs, rewrite first run with OLD: prefix
    for run in demote_p.runs:
        run.text = ""
    if demote_p.runs:
        demote_p.runs[0].text = "OLD: " + full_text
    else:
        demote_p.add_run("OLD: " + full_text)
    print("Demoted to OLD: '%s...'" % full_text[:50])
else:
    print("WARNING: could not find oldest bullet to demote", file=sys.stderr)

# Rotate backups: keep only 3; move older ones to ~/.Trash/
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
while len(all_baks) > 3:
    oldest_bak = all_baks.pop(0)
    trash_dir = os.path.expanduser("~/.Trash/")
    os.makedirs(trash_dir, exist_ok=True)
    dest = os.path.join(trash_dir, os.path.basename(oldest_bak))
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest_bak, dest)
    print("Rotated to Trash: %s" % os.path.basename(oldest_bak))

d.save(DOC)
print("OK: 15:52 backup summary inserted; backup=%s" % os.path.basename(bak))
