#!/usr/bin/env python3
"""Backup ticker 2026-06-28 ~06:07 — update library 'nothing actionable' summary bullet."""
import docx
import sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

NEW_TEXT = (
    "Last tick: 2026-06-28 ~06:07 CDT (backup ticker). "
    "Grovelling apologies — nothing actionable this pass. "
    "Admin queue empty, server healthy (Library + MNU both 200). "
    "All open and in-progress items still awaiting your go-aheads "
    "(sideways-map MNU fix, five-player Cyclopean Gaze crash fix, "
    "MNU missing build files, remote access)."
)

CANDIDATE_PREFIXES = [
    "Last tick: 2026-06-28 ~05:52 CDT (backup ticker)",
    "Last tick: 2026-06-28",
    "Last tick: 2026-06-27",
]

d = docx.Document(DOC)
target = None
for prefix in CANDIDATE_PREFIXES:
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            target = p
            break
    if target is not None:
        break

if target is None:
    print("ERROR: could not find 'Last tick' bullet to update", file=sys.stderr)
    sys.exit(1)

if target.runs:
    target.runs[0].text = NEW_TEXT
    for extra in target.runs[1:]:
        extra.text = ""
else:
    target.add_run(NEW_TEXT)

d.save(DOC)
print(f"OK: updated library 'Last tick' bullet to ~06:07 CDT")
