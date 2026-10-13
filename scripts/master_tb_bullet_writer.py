#!/usr/bin/env python3
import docx
DOC = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx"
NEW = ("[/] Burrowers Beneath (TB) homebrew faction: the guide finished downloading and is now fully "
       "readable. I read it end to end and drafted the entire implementation plan for the faction -- "
       "covering its special \"Mantle\" tile, its unusual multi-part Great Old One (Shudde M'ell's head "
       "plus up to three body segments that move and die separately), all six spellbooks and the "
       "requirements to earn them, the first-to-set-up start with eight tentacles, and the cross-faction "
       "interactions -- and I checked that draft against the general faction-building guide and the "
       "existing factions for consistent menus, overlays, unit sizing and glyphs. The plan is now written "
       "into the TB guide, along with 12 specific design questions for you (for example: exactly when the "
       "\"drop a segment when your Power hits zero\" rule should fire, and how the two enemy-relocation "
       "powers should pick which units move). Please answer those 12 questions inside the TB guide itself "
       "(each one is in red -- reply right after it), and I'll take the faction to its next stage. Nothing "
       "to do in this master list for now.")

d = docx.Document(DOC)
done = False
for p in d.paragraphs:
    if "Burrowers Beneath (TB) guide is now accessible" in p.text:
        for r in list(p.runs):
            r.text = ""
        if p.runs:
            p.runs[0].text = NEW
        else:
            p.add_run(NEW)
        done = True
        break
d.save(DOC)
print("updated bullet 41:", done)
