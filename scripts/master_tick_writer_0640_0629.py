#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~06:40 — update nothing-happened timestamp."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

OLD_PREFIX = "2026-06-29 ~06:33 CDT (master ticker): No new tasks or admin entries."
NEW_TEXT = ("2026-06-29 ~06:40 CDT (master ticker): No new tasks or admin entries. "
            "Both in-progress items still blocked on creator input "
            "(Alternative Ancients clarifications, remote access preference). "
            "Admin queue unreachable (server timeout). Server appears down.")

found = False
for p in d.paragraphs:
    if p.text.strip().startswith(OLD_PREFIX):
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = NEW_TEXT
        else:
            p.add_run(NEW_TEXT)
        print(f"Updated summary bullet timestamp from ~06:33 to ~06:40")
        found = True
        break

if not found:
    print("WARNING: Could not find the ~06:33 summary bullet to update")
    exit(1)

d.save(DOC)
print("OK: saved")
