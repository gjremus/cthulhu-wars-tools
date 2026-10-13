#!/usr/bin/env python3
"""
library_tick_writer_0344_0625_backup.py
Backup ticker safety-net pass, 2026-06-25 ~03:44.
1. Create .bak copy next to docx.
2. Insert new nothing-actionable summary bullet at top of Bulleted summary.
3. Demote the prior ~21:51 nothing-actionable bullet to Old summary entries
   (per Rule 21: MAX 2 nothing-actionable bullets; new one in = oldest one out).
4. Regenerate the .last_seen_tasks.md mirror.
"""
import sys
import shutil
import time
import os
from docx import Document

DOC = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
    "library build current tasks.docx"
)
MIRROR = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/.last_seen_tasks.md"
)

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)
print(f"Backup: {os.path.basename(bak)}")

d = Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~03:44 (backup ticker). Grovelling thanks — overnight "
    "safety-net check: all five games answered healthy and the admin message box is empty. "
    "Nothing changed or deployed; the sideways-map correction and the crashed five-player game "
    "fix both still wait on your go-ahead."
)

ANCHOR_21_51 = "Last tick: 2026-06-24 ~21:51"

# Find the 21:51 anchor paragraph
anchor_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR_21_51):
        anchor_p = p
        break

if anchor_p is None:
    print(f"ERROR: 21:51 summary anchor not found", file=sys.stderr)
    sys.exit(1)

# Insert new summary before the 21:51 paragraph
anchor_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_p.style)
print(f"Inserted new summary before: {anchor_p.text[:60]}...")

# Demote the 21:51 paragraph to just after the Old summary heading
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break

if old_heading is None:
    print("ERROR: Old summary entries heading not found", file=sys.stderr)
    sys.exit(1)

if not anchor_p.text.strip().startswith("OLD:"):
    set_text(anchor_p, "OLD: " + anchor_p.text.strip())
old_heading._p.addnext(anchor_p._p)
print(f"Demoted 21:51 bullet to Old section")

d.save(DOC)
print(f"OK: saved {DOC}")

# Regenerate .last_seen_tasks.md mirror
lines = []
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lines.append("## " + p.text)
    else:
        lines.append(p.text)
with open(MIRROR, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Mirror regenerated: {len(lines)} paragraphs")
