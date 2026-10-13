#!/usr/bin/env python3
"""Append the HB Fix HB-SC-Phase entry to the Homebrew Deployed Online Rollback Guide.docx."""
import docx

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Homebrew Deployed Online Rollback Guide.docx"

d = docx.Document(PATH)

def h1(t): d.add_paragraph(t, style="Heading 1")
def h2(t): d.add_paragraph(t, style="Heading 2")
def body(t): d.add_paragraph(t, style="Body Text")

h1("Fix HB-SC-Phase: Self Consuming stale death tally across phase boundaries (2026-07-07)")

h2("Symptom.")
body("Self Consuming fired at the start of an Action Phase (after Faceless Blight's first move action) and awarded +1 Power even though no units died during that action. The erroneous trigger appeared at the beginning of the third Action Phase in the game 'Failure versus Space'.")

h2("Root cause.")
body("The Self Consuming death counter (fbeSelfConsumingDeaths) accumulates entries whenever any unit is eliminated anywhere in the game via the universal eliminate() hook. The counter is normally cleared at the end of each Action (afterAction). However, unit eliminations that happen during the Doom Phase (e.g., Tablets of the Gods eliminating High Priests during a ritual, or other Doom Phase abilities) or during Power Gather (e.g., prisoner releases from captured-unit return) do NOT go through EndAction/AfterAction — they go through DoomDoneAction or the Power Gather flow instead. These stale deaths persisted into the next Action Phase, causing Self Consuming to fire on the very first action (a harmless move) because afterAction found 2+ entries in the tally left over from a prior phase.")

h2("Fix.")
body("Two one-line additions in solo/Game.scala. (1) At the start of DoomPhaseAction: clear fbeSelfConsumingDeaths = $. (2) At the start of ActionPhaseAction: clear fbeSelfConsumingDeaths = $. This ensures no deaths from Doom Phase or Power Gather leak into the Action Phase, matching the card text ('as part of the same Action' — Doom Phase and Power Gather activities are NOT Actions).")

h2("Files changed.")
body("solo/Game.scala ONLY — two insertions of 'fbeSelfConsumingDeaths = $' (one at DoomPhaseAction, one at ActionPhaseAction).")

h2("Anchor.")
body("solo/Game.scala — case DoomPhaseAction (after doomPhase = true) and case ActionPhaseAction (after fbGhatoLastMoveOrigin = None).")

h2("Revert.")
body("In solo/Game.scala: remove the two 'fbeSelfConsumingDeaths = $' lines from DoomPhaseAction and ActionPhaseAction. Rebuild and redeploy.")

d.save(PATH)
print("Appended Fix HB-SC-Phase entry. Total paragraphs:", len(d.paragraphs))
