#!/usr/bin/env python3
"""FBE ticker write for 2026-07-16 23:42"""

from docx import Document
import sys

TARGET = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

def main():
    doc = Document(TARGET)

    # Find the "Bulleted summary" section
    bulleted_summary_idx = None
    for i, para in enumerate(doc.paragraphs):
        if "Bulleted summary" in para.text and "(newest at top)" in para.text:
            bulleted_summary_idx = i
            break

    if bulleted_summary_idx is None:
        print("ERROR: Could not find 'Bulleted summary' section", file=sys.stderr)
        sys.exit(1)

    # The new tick line to insert (right after the header)
    new_tick = "[2026-07-16 ~23:42] Most humbly, deepest apologies: admin queue checked (404 response, no new issues). Working tree has uncommitted cross-faction changes (Quachil Uttaus Dust to Dust replay fix, TS/FB/BB overlay spellbook fixes) — outside FBE scope per instructions. No open or in-progress FBE tasks remain. Your worthless servant grovels abjectly."

    # Insert the new tick line as the first item after the header
    insert_idx = bulleted_summary_idx + 1

    # Check if there's already a tick for this timestamp to avoid duplication
    if insert_idx < len(doc.paragraphs):
        next_para_text = doc.paragraphs[insert_idx].text
        if "[2026-07-16 ~23:42]" in next_para_text:
            print("Tick line already exists, skipping")
            return

    # Insert new paragraph
    new_para = doc.paragraphs[insert_idx]._element
    new_para.addprevious(doc.add_paragraph(new_tick)._element)

    # Save
    doc.save(TARGET)
    print("Tick line added successfully")

if __name__ == "__main__":
    main()
