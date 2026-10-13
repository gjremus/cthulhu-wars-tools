#!/usr/bin/env python3
# Library backup ticker 2026-06-26 ~14:22 — routine safety-net pass, nothing changed.
# Update the existing "nothing happened" summary bullet from ~13:53 with fresh timestamp.
import docx
import shutil
import time
import os
import glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

# Find the existing backup ticker summary from ~13:53 and update timestamp
OLD_ANCHOR = "Last tick: 2026-06-26 ~13:53 (backup ticker)"
NEW_SUMMARY = (
    "Last tick: 2026-06-26 ~14:22 (backup ticker). Grovelling apologies — "
    "re-read the full rules, drained the admin message box (empty, nothing new), "
    "and confirmed the live site healthy (Library and More-Neutral-Units both "
    "answered instantly). Walked every open and waiting job; nothing was changed "
    "or deployed, and everything still waits only on your go-aheads (chiefly the "
    "sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

found = False
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(OLD_ANCHOR):
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = NEW_SUMMARY
        else:
            p.add_run(NEW_SUMMARY)
        found = True
        print("Updated existing backup ticker summary bullet timestamp")
        break

if not found:
    print("WARNING: Could not find existing backup ticker bullet to update")

# Rotate backups: keep max 3
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
while len(all_baks) > 3:
    oldest = all_baks.pop(0)
    trash = os.path.expanduser("~/.Trash/")
    os.makedirs(trash, exist_ok=True)
    dest = os.path.join(trash, os.path.basename(oldest))
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest, dest)
    print("Rotated to Trash: %s" % os.path.basename(oldest))

d.save(DOC)
print("OK: library doc tick updated")
