#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

# Brief new tick summary (<=3 sentences), grovelling, layman's terms.
NEW_SUMMARY = (
    "Last tick: 2026-06-23 ~17:33. Grovelling apologies — re-read the full rules, found the game "
    "console's message box empty, and walked every open and waiting job; the Library single-player "
    "edition, the More-Neutral-Units beta, the multiplayer game pages and the admin console all "
    "answered healthily. Everything substantial still waits only on your go-aheads (chiefly the "
    "sideways-map correction and the crashed five-player replay); I deployed nothing and changed "
    "nothing on the live site, and the saved copies of this sheet stay at the allowed three."
)

# Brief rewrite of the prior visible summary, honouring the "keep summaries brief" request.
REWRITE_1719 = (
    "Last tick: 2026-06-23 ~17:19. Grovelling apologies — re-read the rules, found the console "
    "message box empty, and confirmed all four live pages healthy. Nothing was deployed; the "
    "substantive jobs still wait only on your go-aheads."
)

ANCHOR_1719 = "Last tick: 2026-06-23 ~17:19"   # newest visible summary
DEMOTE_1703 = "Last tick: 2026-06-23 ~17:03"   # oldest visible nothing-actionable -> OLD section

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


# 1. Insert brief new summary on top (before the 17:19 line).
newest = find_prefix(ANCHOR_1719)
assert newest is not None, "17:19 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. Trim the 17:19 line to its brief form.
set_text(newest, REWRITE_1719)

# 3. Demote the 17:03 line to the Old section (newest-at-top within OLD).
old_para = find_prefix(DEMOTE_1703)
old_head = None
for p in d.paragraphs:
    if p.text.strip().startswith("Old summary entries"):
        old_head = p
        break
demoted = False
if old_para is not None and old_head is not None:
    if old_para.runs:
        old_para.runs[0].text = "OLD: " + old_para.runs[0].text
    else:
        old_para.text = "OLD: " + old_para.text
    old_head._p.addnext(old_para._p)
    demoted = True

# 4. Ensure one empty [] placeholder remains in Open (right before the In progress heading).
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: 17:33 summary added; 17:19 trimmed; demoted 17:03:", demoted,
      "; empty [] present:", has_empty)
