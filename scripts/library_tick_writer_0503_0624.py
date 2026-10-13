#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

# Brief new tick summary (<=3 sentences), grovelling, layman's terms.
NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~05:03. Grovelling apologies for the earlier outage alarm — the live site is "
    "BACK UP: I re-checked from this machine and the Library single-player edition, the More-Neutral-Units "
    "and TchoTcho betas, and the admin message box all answer normally again, and the message queue is "
    "empty. I changed and deployed nothing; everything substantial still waits only on your go-aheads "
    "(chiefly the sideways-map correction and the crashed five-player replay)."
)

ANCHOR_0451 = "Last tick: 2026-06-24 ~04:51"        # current newest summary (the stale DOWN alarm)
SITE_DOWN_INPROG = "[/] INPUT NEEDED — the whole live multiplayer site is DOWN again"
RESOLVED_PREFIX_OLD = "[/] INPUT NEEDED —"
RESOLVED_PREFIX_NEW = (
    "[x] RESOLVED 2026-06-24 ~05:03 — grovelling apologies for the earlier alarm; the live site is back up "
    "and answering normally now (the Library single-player edition, the More-Neutral-Units and TchoTcho "
    "betas, and the admin message box all responded this pass, and the message queue is empty). The outage "
    "has fully cleared and needs nothing further from you; I changed and deployed nothing. "
    "Original alarm kept for history: INPUT NEEDED —"
)

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


# 1. Insert brief new summary on top (before the 04:51 line).
newest = find_prefix(ANCHOR_0451)
assert newest is not None, "04:51 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. Demote the now-stale 04:51 'site DOWN' summary to the Old section (newest-at-top
#    within OLD). The outage it describes has cleared, so it should not sit in the
#    live summary claiming the site is down. Keeps visible summary at the allowed 3.
old_head = None
for p in d.paragraphs:
    if p.text.strip().startswith("Old summary entries"):
        old_head = p
        break
demoted = False
if newest is not None and old_head is not None:
    if newest.runs:
        newest.runs[0].text = "OLD: " + newest.runs[0].text
    else:
        newest.text = "OLD: " + newest.text
    old_head._p.addnext(newest._p)
    demoted = True

# 3. Flip the in-progress 'site DOWN' item to RESOLVED (site recovered this pass).
site = find_prefix(SITE_DOWN_INPROG)
flipped = False
if site is not None and site.runs:
    r0 = site.runs[0]
    if r0.text.startswith(RESOLVED_PREFIX_OLD):
        r0.text = RESOLVED_PREFIX_NEW + r0.text[len(RESOLVED_PREFIX_OLD):]
        flipped = True

# 4. Ensure one empty [] placeholder remains in Open (right before In progress heading).
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: 05:03 summary added; 04:51 demoted:", demoted,
      "; site-down flipped to RESOLVED:", flipped,
      "; empty [] present:", has_empty)
