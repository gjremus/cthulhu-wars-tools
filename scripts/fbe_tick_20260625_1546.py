#!/usr/bin/env python3
"""
FBE tick writer — 2026-06-25 ~15:46
Task: No open tasks. Admin queue empty. Write tick summary (insert only, no archiving).
Currently only 2 summary bullets exist; adding 1 gives exactly 3 (within the ≤3 rule).
"""
import sys
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks.docx"

doc = Document(DOC_PATH)
paras = doc.paragraphs

# Verify the summary section
summary_header_idx = None
for i, p in enumerate(paras):
    if "Bulleted summary" in p.text:
        summary_header_idx = i
        break

if summary_header_idx is None:
    print("ERROR: could not find 'Bulleted summary' paragraph")
    sys.exit(1)

print(f"Found 'Bulleted summary' at index {summary_header_idx}")
for i in range(summary_header_idx, min(summary_header_idx + 6, len(paras))):
    print(f"  [{i}]: {paras[i].text[:120]}")

# Count how many summary bullets we have before the "Open" section
# A summary bullet starts with '[2026-'
bullet_count = 0
first_bullet_idx = None
for i in range(summary_header_idx + 1, len(paras)):
    text = paras[i].text.strip()
    if text.startswith('[2026-'):
        if first_bullet_idx is None:
            first_bullet_idx = i
        bullet_count += 1
    else:
        break  # stop at the first non-bullet (e.g., "Open")

print(f"\nFound {bullet_count} existing summary bullets, first at index {first_bullet_idx}")

if bullet_count >= 3:
    print("ERROR: already at 3 bullets — need to archive first. Aborting.")
    sys.exit(1)

NEW_TICK_LINE = "[2026-06-25 ~15:46] With humble apologies for any delay — this tick found no open Faceless Blight tasks and an empty admin queue; the full implementation audit confirms all six spellbooks, combat, menus, and glyph correct and live as of the last deployment."

from docx.oxml import OxmlElement
from lxml import etree

def clone_para_with_text(source_para, new_text):
    """Clone a paragraph's formatting and replace its text."""
    new_para = deepcopy(source_para._element)
    for r in new_para.findall(qn('w:r')):
        new_para.remove(r)
    r = OxmlElement('w:r')
    source_runs = source_para._element.findall(qn('w:r'))
    if source_runs:
        rpr = source_runs[0].find(qn('w:rPr'))
        if rpr is not None:
            r.append(deepcopy(rpr))
    t = OxmlElement('w:t')
    t.text = new_text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    new_para.append(r)
    return new_para

# Insert new tick bullet BEFORE the first existing bullet
new_tick_elem = clone_para_with_text(paras[first_bullet_idx], NEW_TICK_LINE)
paras[first_bullet_idx]._element.addprevious(new_tick_elem)

# Verify final state
paras_final = doc.paragraphs
summary_header_idx2 = None
for i, p in enumerate(paras_final):
    if "Bulleted summary" in p.text:
        summary_header_idx2 = i
        break

print(f"\nFinal summary section (header at {summary_header_idx2}):")
for i in range(summary_header_idx2, min(summary_header_idx2 + 7, len(paras_final))):
    print(f"  [{i}]: {paras_final[i].text[:120]}")

doc.save(DOC_PATH)
print("\nDOC SAVED.")
