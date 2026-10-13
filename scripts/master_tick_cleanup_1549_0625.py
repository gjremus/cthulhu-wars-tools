#!/usr/bin/env python3
"""Master tick cleanup 2026-06-25 ~15:49:
   Remove orphaned '☒ DONE & LIVE 2026-06-24 ~09:11 (Homebrew game): Bubastis moon-unit sizing'
   paragraph that is stranded in the In-Progress section (it is already in Completed at para 52).
   No marker changes, no task moves beyond this removal.
"""
import docx
from docx.oxml.ns import qn

DOC = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx"

d = docx.Document(DOC)
paras = list(d.paragraphs)
before = len(paras)

# Find the orphaned para: ☒ DONE & LIVE 2026-06-24 ~09:11 in In-Progress section
# It must appear AFTER "In Progress" header and BEFORE "Completed" header
in_progress_idx = None
completed_idx = None
orphan_idx = None

for i, p in enumerate(paras):
    t = p.text.strip()
    if t == "In Progress" and in_progress_idx is None:
        in_progress_idx = i
    if t.startswith("Completed") and completed_idx is None and in_progress_idx is not None:
        completed_idx = i
    if (in_progress_idx is not None and completed_idx is None
            and "2026-06-24 ~09:11 (Homebrew game)" in t
            and t.startswith("☒ DONE")):
        orphan_idx = i

assert in_progress_idx is not None, "Could not find 'In Progress' header"
assert completed_idx is not None, "Could not find 'Completed' header"
assert orphan_idx is not None, (
    f"Could not find orphaned ☒ DONE para between In Progress ({in_progress_idx}) "
    f"and Completed ({completed_idx})"
)

print(f"Found orphaned para at index {orphan_idx}: {paras[orphan_idx].text[:80]}...")

# Remove the paragraph from the XML
p_elem = paras[orphan_idx]._element
p_elem.getparent().remove(p_elem)

after = len(d.paragraphs)
assert after == before - 1, f"paragraph count {before} -> {after} (expected {before - 1})"
d.save(DOC)
print(f"OK: orphaned para removed. paras {before} -> {after}.")
