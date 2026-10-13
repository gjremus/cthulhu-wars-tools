#!/usr/bin/env python3
"""Master ticker 2026-06-30 ~06:29 — update summary, add TT alias task."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# === 1. Update the top summary bullet (replace the ~06:22 backup ticker bullet) ===
in_summary = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary = True
        continue
    if in_summary and p.text.strip().startswith("202"):
        if "06:22 CDT" in p.text and "backup ticker" in p.text:
            for run in p.runs:
                run.text = ""
            p.runs[0].text = (
                "2026-06-30 ~06:29 CDT (master ticker): Both in-progress tasks still "
                "blocked on user input. TT short-link alias broken (players using /TT/ "
                "get a 404; full /TchoTcho/ path works fine) — added as new open task. "
                "Admin queue empty."
            )
            print("OK: updated summary bullet")
            break
        else:
            print(f"INFO: top summary is: {p.text.strip()[:80]}")
            # It may already be a different bullet - insert new one before it
            new_p = OxmlElement('w:p')
            if p._element.find(qn('w:pPr')) is not None:
                new_pPr = deepcopy(p._element.find(qn('w:pPr')))
                new_p.insert(0, new_pPr)
            run_elem = OxmlElement('w:r')
            text_elem = OxmlElement('w:t')
            text_elem.text = (
                "2026-06-30 ~06:29 CDT (master ticker): Both in-progress tasks still "
                "blocked on user input. TT short-link alias broken (players using /TT/ "
                "get a 404; full /TchoTcho/ path works fine) — added as new open task. "
                "Admin queue empty."
            )
            run_elem.append(text_elem)
            new_p.append(run_elem)
            p._element.addprevious(new_p)
            print("OK: inserted new summary bullet before existing top bullet")
            break

# === 2. Add TT alias task to Open section ===
# Find the empty [] bullet in Open and add the new task before it
open_section = False
empty_bullet_para = None
for p in d.paragraphs:
    text = p.text.strip()
    if text == "Open":
        open_section = True
        continue
    if open_section:
        if text in ("In Progress", "Completed"):
            break
        if text == "[]":
            empty_bullet_para = p
            break

if empty_bullet_para:
    # Insert new task before the empty bullet
    new_p = OxmlElement('w:p')
    if empty_bullet_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(empty_bullet_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = (
        "[] TT short-link broken — the /TT/ web address shortcut is returning "
        "'not found' errors. The full /TchoTcho/ address still works, so the "
        "game build itself is fine. The server routing alias needs to be restored."
    )
    run_elem.append(text_elem)
    new_p.append(run_elem)
    empty_bullet_para._element.addprevious(new_p)
    print("OK: added TT alias task to Open")
else:
    print("WARNING: could not find empty [] bullet in Open section")

# === 3. Manage summary bullet count (max 3) ===
# Re-scan summaries
in_summary2 = False
summary_bullets = []
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append(p)
        elif text in ("Open", "In Progress", "Completed") or text.startswith("["):
            break

if len(summary_bullets) > 3:
    # Move excess (oldest = last) to old summary section
    old_summary_para = None
    for p in d.paragraphs:
        if "Old summary bullets" in p.text:
            old_summary_para = p
            break
    if old_summary_para:
        while len(summary_bullets) > 3:
            oldest = summary_bullets.pop()
            oldest_text = oldest.text.strip()
            # Copy to old summaries
            new_p = OxmlElement('w:p')
            if oldest._element.find(qn('w:pPr')) is not None:
                new_pPr = deepcopy(oldest._element.find(qn('w:pPr')))
                new_p.insert(0, new_pPr)
            run_elem = OxmlElement('w:r')
            text_elem = OxmlElement('w:t')
            text_elem.text = oldest_text
            run_elem.append(text_elem)
            new_p.append(run_elem)
            old_summary_para._element.addnext(new_p)
            # Remove from summary
            oldest._element.getparent().remove(oldest._element)
            print(f"OK: moved oldest summary to old-summaries")

d.save(DOC)
print("SAVED")
