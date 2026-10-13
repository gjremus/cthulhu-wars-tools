#!/usr/bin/env python3
"""Move server-crash task from Open to In Progress section (rule 10: [/] items go in In Progress)."""
import sys
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

doc = Document(DOC_PATH)
paras = doc.paragraphs

# Find the server-crash task paragraph
TASK_TEXT = "INPUT NEEDED — Server-side crash (not our code)"
task_idx = None
for i, p in enumerate(paras):
    if TASK_TEXT in p.text:
        task_idx = i
        break

if task_idx is None:
    print("ERROR: Could not find server-crash task")
    sys.exit(1)

print(f"Found task at index {task_idx}: {paras[task_idx].text[:80]}")

# Find "In Progress" section header
inprog_idx = None
for i, p in enumerate(paras):
    if p.text.strip() == 'In Progress':
        inprog_idx = i
        break

if inprog_idx is None:
    print("ERROR: Could not find 'In Progress' header")
    sys.exit(1)

print(f"Found 'In Progress' at index {inprog_idx}")

# The task is currently in Open (before In Progress).
# We need to remove it from Open and insert it after the In Progress header
# (before the existing [/] task there).

# Save the element reference
task_elem = paras[task_idx]._element

# Remove from current position
task_elem.getparent().remove(task_elem)

# Find the first item after "In Progress" header in the refreshed doc
# After removal, indices shift. Let's re-find "In Progress"
paras2 = doc.paragraphs
for i, p in enumerate(paras2):
    if p.text.strip() == 'In Progress':
        # Insert the task element right after this heading
        # The next paragraph should be the existing [/] deploy task
        if i + 1 < len(paras2):
            paras2[i+1]._element.addprevious(task_elem)
            print(f"Moved task to In Progress section (before index {i+1})")
        else:
            # Append after In Progress heading
            p._element.addnext(task_elem)
            print("Appended task after In Progress heading")
        break

doc.save(DOC_PATH)
print("DOC SAVED.")
