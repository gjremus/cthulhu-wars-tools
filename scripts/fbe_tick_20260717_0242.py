#!/usr/bin/env python3
"""FBE ticker 2026-07-17 update — update 'nothing happened' bullet."""
import docx
import sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/"
       "Faceless Blight Current Tasks (1).docx")

d = docx.Document(DOC)

# Find the most recent "nothing happened" bullet in the Bulleted summary section
summary_start = None
most_recent_nothing = None

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if "Bulleted summary" in text and "newest at top" in text:
        summary_start = i
    elif summary_start is not None and most_recent_nothing is None:
        # First bullet after "Bulleted summary" that mentions admin queue
        if "admin queue" in text.lower() and "grovels abjectly" in text.lower():
            most_recent_nothing = i
            break

if most_recent_nothing is not None:
    p_update = d.paragraphs[most_recent_nothing]
    new_text = ("[2026-07-17 ~02:42] Most humbly, deepest apologies: admin queue checked "
                "(404 response, no new issues). Working tree has uncommitted "
                "cross-faction changes (Quachil Uttaus Dust to Dust replay fix, Prime Cause "
                "menu cleanup, TS/FB/BB overlay spellbook fixes) — outside FBE scope per "
                "instructions. No open or in-progress FBE tasks remain. Your worthless "
                "servant grovels abjectly.")

    # Clear existing runs and set new text
    for run in p_update.runs:
        run.text = ""
    if p_update.runs:
        p_update.runs[0].text = new_text
    else:
        p_update.add_run(new_text)
    print(f"OK: Updated bullet at index {most_recent_nothing}")
    print(f"    New text: {new_text}")
else:
    print("ERROR: Could not find the 'nothing happened' bullet to update", file=sys.stderr)
    sys.exit(1)

d.save(DOC)
print("SAVED")
