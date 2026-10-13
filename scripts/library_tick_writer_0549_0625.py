#!/usr/bin/env python3
"""
library_tick_writer_0549_0625.py
Library ticker pass, 2026-06-25 ~05:49 CDT.
1. Create .bak copy next to docx.
2. Insert new summary bullet at top of Bulleted summary.
3. Demote the "Last tick: 2026-06-25 ~03:44 (backup ticker)" nothing-actionable
   bullet to Old summary entries (per Rule 21: adding a new nothing-actionable
   bullet when 2 already exist means oldest nothing-actionable moves out).
4. Regenerate the .last_seen_tasks.md mirror.
No task status changes this pass — all INPUT NEEDED items still waiting,
no @@@ responses from user detected.
"""
import sys
import shutil
import time
import os
from docx import Document

DOC = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
    "library build current tasks.docx"
)
MIRROR = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/.last_seen_tasks.md"
)

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)
print(f"Backup: {os.path.basename(bak)}")

d = Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~05:49. Grovelling thanks — re-read the full rules, drained "
    "the admin message box (empty, nothing new), confirmed all live editions healthy "
    "(the Library single-player edition, the More-Neutral-Units beta, and the known-good "
    "game page all answered instantly), and walked every open and waiting job; nothing was "
    "changed or deployed, and everything still waits only on your go-aheads (chiefly the "
    "sideways-map correction and the crashed five-player game fix)."
)

# Anchor: the ~03:44 backup-ticker nothing-actionable bullet to demote to OLD
ANCHOR_0344 = "Last tick: 2026-06-25 ~03:44"

# ── Step 1: Find the ~05:34 paragraph (the current top bullet) ───────────────
anchor_0534_p = None
for p in d.paragraphs:
    if p.text.strip().startswith("Last tick: 2026-06-25 ~05:34"):
        anchor_0534_p = p
        break

if anchor_0534_p is None:
    print("ERROR: 05:34 summary anchor not found", file=sys.stderr)
    sys.exit(1)

# ── Step 2: Insert new summary before the ~05:34 paragraph ───────────────────
anchor_0534_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_0534_p.style)
print("Inserted new summary before 05:34 bullet")

# ── Step 3: Find the ~03:44 paragraph ────────────────────────────────────────
anchor_0344_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR_0344):
        anchor_0344_p = p
        break

if anchor_0344_p is None:
    print(f"ERROR: 03:44 summary anchor not found", file=sys.stderr)
    sys.exit(1)

# ── Step 4: Demote the ~03:44 paragraph to Old summary entries ───────────────
old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break

if old_heading is None:
    print("ERROR: Old summary entries heading not found", file=sys.stderr)
    sys.exit(1)

if not anchor_0344_p.text.strip().startswith("OLD:"):
    set_text(anchor_0344_p, "OLD: " + anchor_0344_p.text.strip())
old_heading._p.addnext(anchor_0344_p._p)
print("Demoted 03:44 bullet to Old section")

# ── Step 5: Save ─────────────────────────────────────────────────────────────
d.save(DOC)
print(f"OK: saved {DOC}")

# ── Step 6: Regenerate .last_seen_tasks.md mirror ────────────────────────────
lines = []
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lines.append("## " + p.text)
    else:
        lines.append(p.text)
with open(MIRROR, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Mirror regenerated: {len(lines)} paragraphs")
