#!/usr/bin/env python3
"""Master ticker 2026-06-30 ~23:26 — consolidate 'nothing happened' bullets."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# Find key paragraphs containing the two "nothing actionable" bullets
nothing_indices = []
for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if "Nothing actionable" in text or "Nothing actionable" in text.lower():
        nothing_indices.append(i)
    elif "nothing actionable" in text.lower():
        nothing_indices.append(i)

print(f"Found 'nothing actionable' paragraphs at indices: {nothing_indices}")
for idx in nothing_indices:
    print(f"  [{idx}] {d.paragraphs[idx].text.strip()[:100]}")

# Strategy: Update the FIRST one with current timestamp, clear the SECOND one
if len(nothing_indices) >= 2:
    # Update first nothing-happened bullet with current timestamp
    p_keep = d.paragraphs[nothing_indices[0]]
    new_text = ("2026-06-30 ~23:26 CDT (master ticker): Admin queue empty. "
                "All 3 in-progress tasks still blocked on user input "
                "(Alt Ancients awaiting Section 1 answers, Telegram awaiting bot token, "
                "BB game 540 awaiting decision). Nothing actionable.")
    for run in p_keep.runs:
        run.text = ""
    if p_keep.runs:
        p_keep.runs[0].text = new_text
    else:
        p_keep.add_run(new_text)
    print(f"OK: Updated bullet at index {nothing_indices[0]}")

    # Remove the second one by clearing its text
    p_remove = d.paragraphs[nothing_indices[1]]
    for run in p_remove.runs:
        run.text = ""
    # Remove the paragraph element entirely from the XML
    p_remove._element.getparent().remove(p_remove._element)
    print(f"OK: Removed bullet at index {nothing_indices[1]}")

elif len(nothing_indices) == 1:
    # Just update the one we have
    p_keep = d.paragraphs[nothing_indices[0]]
    new_text = ("2026-06-30 ~23:26 CDT (master ticker): Admin queue empty. "
                "All 3 in-progress tasks still blocked on user input "
                "(Alt Ancients awaiting Section 1 answers, Telegram awaiting bot token, "
                "BB game 540 awaiting decision). Nothing actionable.")
    for run in p_keep.runs:
        run.text = ""
    if p_keep.runs:
        p_keep.runs[0].text = new_text
    else:
        p_keep.add_run(new_text)
    print(f"OK: Updated single bullet at index {nothing_indices[0]}")
else:
    print("No 'nothing actionable' bullets found — unexpected state")

d.save(DOC)
print("SAVED")
