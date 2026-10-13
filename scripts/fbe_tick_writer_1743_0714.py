#!/usr/bin/env python3
"""FBE ticker 2026-07-14 ~17:43 — no work, updating existing tick."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/"
       "Faceless Blight Current Tasks (1).docx")

d = docx.Document(DOC)

# Find the most recent tick line (should be first in Bulleted summary section)
# Looking for the line starting with [2026-07-14 ~17:27]
target_idx = None
for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text.startswith("[2026-07-14 ~17:27]"):
        target_idx = i
        print(f"Found target paragraph at index {i}")
        break

if target_idx is not None:
    p = d.paragraphs[target_idx]
    new_text = ("[2026-07-14 ~17:43] Most humbly, deepest apologies: admin queue empty "
                "(no FBE-specific issues). No open or in-progress tasks. "
                "Your worthless servant grovels abjectly.")

    # Clear existing runs and set new text
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

    print(f"OK: Updated tick line at index {target_idx}")
else:
    print("ERROR: Could not find target paragraph starting with [2026-07-14 ~17:27]")
    exit(1)

d.save(DOC)
print("SAVED")
