#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~18:33 — routine healthy pass, nothing changed/deployed.
# Current visible bullets: [17:52, 15:39]. Both are nothing-actionable.
# Rule 22: max 2 nothing-actionable bullets. Adding mine makes 3 → demote 15:39 to OLD.
import docx
import shutil
import time
import os
import sys
import glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~18:33 (backup ticker). Grovelling apologies — re-read "
    "the full rules, drained the admin message box (empty, nothing new), and confirmed "
    "all five games fully healthy (Library, More-Neutral-Units, TchoTcho, Bubastis, and "
    "Homebrew all answered HTTP 200). Walked every open and waiting job; nothing was "
    "changed or deployed, and everything still waits only on your go-aheads (chiefly "
    "the sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

ANCHOR = "Last tick: 2026-06-25 ~17:52"
DEMOTE_ANCHOR = "Last tick: 2026-06-25 ~15:39"

anchor_p = None
demote_p = None

for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(ANCHOR) and not t.startswith("OLD:"):
        anchor_p = p
    if t.startswith(DEMOTE_ANCHOR) and not t.startswith("OLD:"):
        demote_p = p

if anchor_p is None:
    print("ERROR: could not find anchor paragraph starting with '%s'" % ANCHOR, file=sys.stderr)
    sys.exit(1)

# Insert new summary bullet before the anchor (newest at top)
anchor_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_p.style)
print("Inserted new 18:33 summary before anchor.")

# Demote the oldest nothing-actionable bullet to OLD: per rule 22
if demote_p is not None:
    full_text = demote_p.text
    for run in demote_p.runs:
        run.text = ""
    if demote_p.runs:
        demote_p.runs[0].text = "OLD: " + full_text
    else:
        demote_p.add_run("OLD: " + full_text)
    print("Demoted to OLD: '%s...'" % full_text[:60])
else:
    print("WARNING: could not find oldest bullet to demote", file=sys.stderr)

# Rotate backups: keep only 3; move older ones to ~/.Trash/
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
while len(all_baks) > 3:
    oldest_bak = all_baks.pop(0)
    trash_dir = os.path.expanduser("~/.Trash/")
    os.makedirs(trash_dir, exist_ok=True)
    dest = os.path.join(trash_dir, os.path.basename(oldest_bak))
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest_bak, dest)
    print("Rotated to Trash: %s" % os.path.basename(oldest_bak))

d.save(DOC)
print("OK: 18:33 backup summary inserted; backup=%s" % os.path.basename(bak))
