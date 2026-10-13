#!/usr/bin/env python3
"""Append the HB Fix 115 entry to the Homebrew Deployed Online Rollback Guide.docx,
matching the existing Heading1 / Heading2 + Body Text structure of prior fixes."""
import docx

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/HomeBrews/Homebrew Deployed Online Rollback Guide.docx"

d = docx.Document(PATH)

def h1(t): d.add_paragraph(t, style="Heading 1")
def h2(t): d.add_paragraph(t, style="Heading 2")
def body(t): d.add_paragraph(t, style="Body Text")

h1("Fix 115 (2026-06-14) — DC Tenebrosum now fires after awakening an INDEPENDENT GOO")

h2("Symptom.")
body("After Defilers Court awakened an independent Great Old One (iGOO), the conditional unlimited “Tenebrosum: Repeat Awaken for N Sin” button never appeared — so DC could not pay Sin to awaken a second (different) GOO/iGOO that same turn, even when one was still awakenable. Awakening a faction GOO (Y’Golonac) via the normal Awaken action worked fine; only iGOO awakens failed to offer the Tenebrosum repeat.")

h2("Root cause.")
body("The Tenebrosum “last action” recorder in solo/FactionDC.scala only matched AwakenAction. Awakening an iGOO does NOT produce an AwakenAction — it produces IndependentGOOAction (solo/IGOOs.scala:216), and Cthugha produces CthughaAwakenAction (IGOOs.scala:279). Both extend BaseFactionAction directly and are NOT subtypes of AwakenAction, so recordTenebrosum() was never called, game.dcLastActionForTenebrosum stayed None, and the offer site (wrapped in dcLastActionForTenebrosum.foreach) rendered nothing. The Fix 113 broadened-Awaken eligibility check (hasPoolGOO || hasIGOO) was correct but never reached because nothing was recorded.")

h2("Fix.")
body("Three surgical edits, all in solo/FactionDC.scala. (1) Added recording cases for IndependentGOOAction(self,_,_,cost) and CthughaAwakenAction(self,_,_,cost) right after the AwakenAction case — each calls recordTenebrosum(action, cost, “Awaken”) using the action’s own carried cost. (2) Broadened the tenebrosumLegalToRepeat case label from `case _ : AwakenAction` to `case _ : AwakenAction | _ : IndependentGOOAction | _ : CthughaAwakenAction` so iGOO awakens use the same hasPoolGOO || hasIGOO eligibility. (3) Broadened the tenebrosumRepeatChooser case label identically so iGOO awakens route through DCTenebrosumAwakenChooserAction — the broadened chooser that offers EVERY awakenable faction GOO + iGOO (consistent with Fix 113), not just the exact iGOO just awakened.")

h2("Files changed.")
body("solo/FactionDC.scala ONLY — two new recording cases plus two broadened case labels (HB Fix 115 in-line comment). No other source touched; IndependentGOOAction / CthughaAwakenAction are public case classes already visible in this file.")
body("NOT included in this build (pre-existing, in-progress, does NOT compile — same as Fix 113/114): the in-flight Faceless Blight (FBE) faction WIP across solo/Battle.scala, solo/BotFBE.scala, solo/FactionFBE.scala, solo/Game.scala. These were git-stashed before building so the shipped HB build contains ONLY the Fix 115 change on top of Fix 114, then restored to the working tree afterward (owner WIP untouched).")
body("Version NOT bumped: common.sbt still reads Homebrew v 2.6 (this fix shipped without a version-string bump to avoid an unnecessary second production redeploy; the bump will ride the next HB build that bundles the FBE work).")

h2("Build + ship.")
body("Stashed the 4 FBE-WIP files, then: cd /Users/gremus/Claude-Projects/cw-homebrew-wt/server-deploy && ./deploy-homebrew-to-vm.sh --build (sbt fullOptJS). Closure: 0 error(s), 0 warning(s). Cache tag 20260614-183208, public HEAD HTTP 200 at https://cwo.freeddns.org/HB/. Restored the FBE WIP to the working tree after deploy.")

h2("Git commit.")
body("011faf0 on branch homebrew (repo gjremus/cthulhu-wars), pushed to origin/homebrew (3a433f8..011faf0). Message: “HB Fix 115: DC Tenebrosum now fires after awakening an independent GOO”. Touches exactly solo/FactionDC.scala.")

h2("Rollback.")
body("git revert 011faf0 (touches only solo/FactionDC.scala), rebuild Homebrew (sbt fullOptJS, with FBE WIP stashed), redeploy. Restores Fix 114 behavior — i.e. the Tenebrosum repeat button stops appearing after an iGOO awaken.")

d.save(PATH)
print("Appended Fix 115 entry. Total paragraphs:", len(d.paragraphs))
