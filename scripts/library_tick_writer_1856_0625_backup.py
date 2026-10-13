#!/usr/bin/env python3
# Library backup ticker 2026-06-25 ~18:56 — routine healthy pass, nothing changed/deployed.
# Current visible bullets: [18:33, 17:52 (demoted to OLD by 18:33 writer)].
# Actually looking at last_seen, the current visible ones are 18:33 and the one it kept.
# Rule 22: max 2 nothing-actionable bullets. Replace the 18:33 bullet with updated time.
import docx
import shutil
import time
import os
import sys
import glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

# Find the most recent summary bullet and update its timestamp in-place
# This avoids accumulating bullets — just updates the time per rule 22
OLD_ANCHOR = "Last tick: 2026-06-25 ~18:33"
NEW_TEXT = (
    "Last tick: 2026-06-25 ~18:56 (backup ticker). Grovelling apologies — re-read "
    "the full rules, drained the admin message box (empty, nothing new), and confirmed "
    "all five games fully healthy (Library, More-Neutral-Units, TchoTcho, Bubastis, and "
    "Homebrew all answered HTTP 200). Walked every open and waiting job; nothing was "
    "changed or deployed, and everything still waits only on your go-aheads (chiefly "
    "the sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

found = False
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(OLD_ANCHOR) and not t.startswith("OLD:"):
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = NEW_TEXT
        else:
            p.add_run(NEW_TEXT)
        found = True
        print("Updated library summary from ~18:33 to ~18:56")
        break

if not found:
    print("ERROR: could not find anchor '%s'" % OLD_ANCHOR, file=sys.stderr)
    sys.exit(1)

# Rotate backups: keep only 3
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
while len(all_baks) > 3:
    oldest_bak = all_baks.pop(0)
    trash_dir = os.path.expanduser("~/.Trash/")
    os.makedirs(trash_dir, exist_ok=True)
    dest = os.path.join(trash_dir, os.path.basename(oldest_bak))
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest_bak, dest)
    print("Rotated to Trash: %s" % os.path.basename(oldest_bak))

d.save(DOC)
print("OK: library backup tick updated; backup=%s" % os.path.basename(bak))
