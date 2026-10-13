#!/usr/bin/env python3
"""FBE ticker 2026-07-18 ~11:06 — Shapestealing dice roll fix completed."""
from docx import Document

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx")

doc = Document(DOC)

# Step 1: Remove the triple-at marked input from Open section
in_open = False
open_marker_removed = False
paragraphs_to_remove = []

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text == "Open":
        in_open = True
        continue
    if in_open and "Claude Questions for the Creator" in text:
        in_open = False
        break
    if in_open and "@@@" in p.text and "Shapestealing" in p.text:
        # Mark for removal (need to remove both the marker and following lines)
        paragraphs_to_remove.append(i)
        # Check if next few paragraphs are continuation
        for j in range(i+1, min(i+5, len(doc.paragraphs))):
            next_text = doc.paragraphs[j].text.strip()
            if next_text.startswith("-") or next_text == "" or "Claude Questions" in next_text:
                break
            paragraphs_to_remove.append(j)
        open_marker_removed = True
        break

# Remove paragraphs in reverse order to avoid index shifting
for idx in reversed(paragraphs_to_remove):
    p = doc.paragraphs[idx]
    p_element = p._element
    p_element.getparent().remove(p_element)

if open_marker_removed:
    print(f"OK: removed {len(paragraphs_to_remove)} paragraph(s) from Open section")
else:
    print("WARNING: could not find triple-at input in Open section")

# Step 2: Add tick summary at top of Bulleted summary
found_summary = False
insertion_index = None

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if "Bulleted summary" in text:
        # Insert after the heading
        if i + 1 < len(doc.paragraphs):
            next_text = doc.paragraphs[i + 1].text.strip()
            if next_text == "" or "(newest at top)" in next_text:
                insertion_index = i + 2
            else:
                insertion_index = i + 1
        else:
            insertion_index = i + 1
        found_summary = True
        break

if not found_summary:
    print("ERROR: could not find Bulleted summary section")
    exit(1)

new_para = doc.paragraphs[insertion_index].insert_paragraph_before()
new_para.text = ("[2026-07-18 ~11:06] Most humbly, deepest apologies: Shapestealing fix completed — "
                 "stolen monsters now fight for your side only (was incorrectly rolling for both sides). "
                 "Commit a3f38e6, ready for build and deploy. Your worthless servant grovels abjectly.")
print("OK: added tick summary")

# Step 3: Add to Completed section
in_completed = False
completed_insertion_index = None

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text == "Completed":
        in_completed = True
        # Find first bullet in Completed section
        for j in range(i+1, len(doc.paragraphs)):
            next_text = doc.paragraphs[j].text.strip()
            if next_text.startswith("-") or next_text.startswith("•"):
                completed_insertion_index = j
                break
        break

if completed_insertion_index is None:
    print("ERROR: could not find Completed section or first bullet")
    exit(1)

new_completed = doc.paragraphs[completed_insertion_index].insert_paragraph_before()
new_completed.text = ("- ☒ Shapestealing dice roll fix — FIXED (2026-07-18). When a monster was successfully "
                      "stolen via Shapestealing, both sides were incorrectly rolling combat dice for it. "
                      "Root cause: the battle strength calculator was looking at the wrong unit list after "
                      "the stolen monster was transferred to your side. Fixed so only the side that controls "
                      "the stolen monster rolls dice for it. Commit a3f38e6, ready for build and deploy. "
                      "Your worthless servant grovels abjectly.")
print("OK: added to Completed section")

doc.save(DOC)
print("SAVED")
