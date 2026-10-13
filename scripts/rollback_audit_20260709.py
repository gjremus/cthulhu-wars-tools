#!/usr/bin/env python3
"""Append missing rollback guide entries for fixes deployed July 1-9, 2026."""
import docx

GUIDES = {
    "bb": "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/Bubastis/Bubastis Deployed Online Rollback Guide.docx",
    "hb": "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Homebrew Deployed Online Rollback Guide.docx",
    "mnu": "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Neutral Units/More Neutral Units/MNU Rollback Guide.docx",
    "library": "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/Library Deployed Online Rollback Guide.docx",
}

def append_entries(path, entries):
    d = docx.Document(path)
    def h1(t): d.add_paragraph(t, style="Heading 1")
    def h2(t): d.add_paragraph(t, style="Heading 2")
    def body(t): d.add_paragraph(t, style="Body Text")

    for entry in entries:
        h1(entry["title"])
        for section in entry["sections"]:
            h2(section[0])
            body(section[1])

    d.save(path)
    print(f"  Appended {len(entries)} entries. Total paragraphs: {len(d.paragraphs)}")


# === BB ENTRIES ===
bb_entries = [
    {
        "title": "Fix BB-Pain: Lingering pain indicator on units after battle (2026-07-09)",
        "sections": [
            ("Symptom.", "Units that participated in battle (specifically IGOOs like Y'Golonac) could retain a visual Pained health state permanently, showing red pain scratch lines on the map even after battle ended and no combat was active."),
            ("Root cause.", "The BattleEnd health reset (sides.foreach(_.forces.foreach(u => u.health = Alive))) only covers units still in the `forces` list. Units that left forces during battle (e.g., IGOOs via exemption, special handling, or death-avoidance mechanics) retained their Pained state permanently because no other code path reset it."),
            ("Fix.", "Added a global safety net after the existing forces-based reset in Battle.scala BattleEnd: game.factions.foreach(_.units.foreach(u => if (u.health == Pained) u.health = Alive)). This catches ANY unit across ALL factions that still has Pained health after battle resolution."),
            ("Anchor.", "solo/Battle.scala — BattleEnd phase, immediately after the three sides.foreach lines (forces health reset, Retreated removal, Zeroed removal). Line ~1757 in the BB build."),
            ("Revert.", "Remove the game.factions.foreach(_.units.foreach(u => if (u.health == Pained) u.health = Alive)) line from Battle.scala BattleEnd. Rebuild and redeploy."),
        ]
    },
]

# === HB ENTRIES ===
hb_entries = [
    {
        "title": "Fix HB-DarkBargain: DC Dark Bargain log visibility for multiplayer (2026-07-08)",
        "sections": [
            ("Symptom.", "In multiplayer DC games, when a player confirmed a Dark Bargain choice, the log entry didn't appear until after the next action, leaving opponents unable to see what was chosen in real-time."),
            ("Root cause.", "The log call was inside the resolution handler rather than immediately upon confirm."),
            ("Fix.", "Moved the log statement to fire immediately on confirm action in the DC Dark Bargain handler."),
            ("Anchor.", "solo/Game.scala — DC Dark Bargain confirm action handler."),
            ("Revert.", "Revert commit 0988b4b on homebrew branch. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix HB-Satiate: DC Satiate multi-capture + gate diplomacy + Holy Ground (2026-07-08)",
        "sections": [
            ("Symptom.", "Multiple bugs in DC Satiate: (1) multi-capture didn't properly trigger FBE Self Consuming; (2) gate diplomacy wasn't checked during Satiate captures; (3) Holy Ground could block captures that should be allowed."),
            ("Root cause.", "Satiate's capture loop wasn't routing through the standard capture pathway that handles these edge cases."),
            ("Fix.", "Three commits: ba616ba (Self Consuming proc), 655cdbf (gate diplomacy + Holy Ground), 9b26847 (skip prompt for equivalent cultists)."),
            ("Anchor.", "solo/Game.scala — DC Satiate action handlers."),
            ("Revert.", "Revert commits ba616ba, 655cdbf, and 9b26847 on homebrew branch. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix HB-Pain: Lingering pain indicator on units after battle (2026-07-09)",
        "sections": [
            ("Symptom.", "Units that left the battle forces list (e.g., IGOOs) could retain Pained health state permanently, showing red pain scratch lines on the map after battle ended."),
            ("Root cause.", "BattleEnd health reset only covers units in the `forces` list. Units that left forces during battle retained Pained state."),
            ("Fix.", "Added global safety net at BattleEnd: game.factions.foreach(_.units.foreach(u => if (u.health == Pained) u.health = Alive))."),
            ("Anchor.", "solo/Battle.scala — BattleEnd phase, after the three sides.foreach lines. Line ~1964 in the HB build."),
            ("Revert.", "Remove the game.factions.foreach line from Battle.scala BattleEnd. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix HB-DS-Display: DS alternate spellbooks showing as empty ? slots (2026-07-09)",
        "sections": [
            ("Symptom.", "When DSAlternateSpellbooks option is active, DS faction card shows empty '?' spellbook slots for books that were actually earned (Omnipotence, DirectedEnergy, FiendishSpawn)."),
            ("Root cause.", "The display formula f.spellbooks.%(f.library.has).num only counted books literally in the original library list. Alternate spellbooks (Omnipotence replacing Traitors, etc.) weren't counted, making 2+ earned slots appear as unfilled '?' marks."),
            ("Fix.", "Compute an effectiveLibrary that maps alternates (same logic as CheckSpellbooksAction), then use effectiveLib.has instead of f.library.has for the slot count. Also handles SleeperEnergyNexusPreBattle alternate."),
            ("Anchor.", "solo/CthulhuWarsSolo.scala — faction card spellbook display block, the '?' rendering line (around line 2714 in HB build)."),
            ("Revert.", "Replace the effectiveLib block with the original single line: (1.to(f.library.num - f.spellbooks.%(f.library.has).num - f.unfulfilled.num)... Rebuild and redeploy."),
        ]
    },
]

# === MNU ENTRIES ===
mnu_entries = [
    {
        "title": "Fix MNU-ArcticWind: WW Arctic Wind onGate flag not clearing (2026-07-08)",
        "sections": [
            ("Symptom.", "After Wendigo's Arctic Wind forced-follow moved units, their onGate flag remained set, causing them to appear as if they were still on a gate in the moved-to region."),
            ("Root cause.", "The Arctic Wind forced-follow handler didn't clear the onGate unit flag after the move completed."),
            ("Fix.", "Added onGate flag clear in the Arctic Wind forced-follow handler."),
            ("Anchor.", "solo/FactionWW.scala — Arctic Wind forced-follow action handler."),
            ("Revert.", "Revert commit ffb5d6b on more-neutral-units branch. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix MNU-Pain: Lingering pain indicator on units after battle (2026-07-09)",
        "sections": [
            ("Symptom.", "Units that left the battle forces list could retain Pained health state permanently after battle ended."),
            ("Root cause.", "BattleEnd health reset only covers units in the `forces` list."),
            ("Fix.", "Added global safety net at BattleEnd: game.factions.foreach(_.units.foreach(u => if (u.health == Pained) u.health = Alive))."),
            ("Anchor.", "solo/Battle.scala — BattleEnd phase, after the three sides.foreach lines. Line ~1577 in the MNU build."),
            ("Revert.", "Remove the game.factions.foreach line from Battle.scala BattleEnd. Rebuild and redeploy."),
        ]
    },
]

# === LIBRARY ENTRIES ===
library_entries = [
    {
        "title": "Fix Library-CG-SBR: CG/SBR replay divergence (2026-07-01)",
        "sections": [
            ("Symptom.", "Games with Cyclopean Gaze (CG) could diverge on replay if a spellbook requirement was satisfied at the same time CG fired, because CheckSpellbooksAction wasn't wrapping the CG action."),
            ("Root cause.", "CG was called via Force() instead of CheckSpellbooksAction(), so any SBR satisfied in the same afterAction would not prompt the spellbook choice before CG executed."),
            ("Fix.", "Wrapped CG fire action in CheckSpellbooksAction. Also added afterAction SBR checks for OW and BG before CG fires."),
            ("Anchor.", "solo/Game.scala — AfterAction section, CG fire logic."),
            ("Revert.", "Revert commits 996c020 and 9748a4d on library-at-celaeno branch. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix Library-ENPB: Energy Nexus Pre-Battle resume applying to wrong battle (2026-07-02)",
        "sections": [
            ("Symptom.", "When Energy Nexus Pre-Battle variant fired during a Grasping Dead battle, the game gave Sleeper a full turn but then failed to return to the interrupted battle, or returned to the wrong one."),
            ("Root cause.", "NextPlayerAction routing didn't account for nexed/battleResumePhase state, so it continued the normal turn order instead of resuming the interrupted battle."),
            ("Fix.", "Added nexed/battleResumePhase checks to NextPlayerAction routing and added a resume block to ProceedBattlesAction."),
            ("Anchor.", "solo/Game.scala — NextPlayerAction routing logic + ProceedBattlesAction entry block."),
            ("Revert.", "Revert commit 63a0688 on library-at-celaeno branch. Rebuild and redeploy."),
        ]
    },
    {
        "title": "Fix Library-Pain: Lingering pain indicator on units after battle (2026-07-09)",
        "sections": [
            ("Symptom.", "Units that left the battle forces list could retain Pained health state permanently after battle ended."),
            ("Root cause.", "BattleEnd health reset only covers units in the `forces` list."),
            ("Fix.", "Added global safety net at BattleEnd: game.factions.foreach(_.units.foreach(u => if (u.health == Pained) u.health = Alive))."),
            ("Anchor.", "solo/Battle.scala — BattleEnd phase, after the three sides.foreach lines. Line ~1114 in the Library build."),
            ("Revert.", "Remove the game.factions.foreach line from Battle.scala BattleEnd. Rebuild and redeploy."),
        ]
    },
]

# === WRITE ALL ===
import sys

target = sys.argv[1] if len(sys.argv) > 1 else "all"

if target in ("bb", "all"):
    print("BB guide:")
    append_entries(GUIDES["bb"], bb_entries)

if target in ("hb", "all"):
    print("HB guide:")
    append_entries(GUIDES["hb"], hb_entries)

if target in ("mnu", "all"):
    print("MNU guide:")
    append_entries(GUIDES["mnu"], mnu_entries)

if target in ("library", "all"):
    print("Library guide:")
    append_entries(GUIDES["library"], library_entries)

print("Done.")
