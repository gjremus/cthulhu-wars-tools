#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

NEW = (
    "Last tick: 2026-06-15 10:05 (next ~10:20). Grovelling apologies — re-read the rules, drained the "
    "game-console queue (only the three already-known notes: the crashed five-player practice game and "
    "the master helper's dice-roll hand-off) and walked every open and in-progress job; the main site, "
    "the More-Neutral-Units edition and the admin console all answered healthily and the saved copies of "
    "this sheet stay at three. Genuinely new this pass: the 'missing publishing recipe' is recoverable — "
    "I found the sister edition's working publish-to-the-live-site recipe still on this machine and "
    "confirmed the Library edition's own up-to-date source is present too, so publishing the finished "
    "dice-roll fix is now a concrete go/no-go rather than a mystery. It still needs your nod on the exact "
    "method (let me carefully adapt that sister recipe and push, backing up the live copy first), and I "
    "still need your OK to also correct the sideways map and the related game-freeze on the edition you "
    "actually see."
)

DEMOTE_PREFIX = "Last tick: 2026-06-15 09:18"  # oldest of the current 3 -> move to Old section

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def find_heading_contains(text):
    for p in d.paragraphs:
        if text in p.text.strip() and p.style.name.startswith("Heading"):
            return p
    return None


newest = find_prefix("Last tick: 2026-06-15 09:50")
assert newest is not None, "09:50 summary anchor missing"

# 1. New brief summary on top.
newest.insert_paragraph_before(NEW, style=newest.style)

# 2. Demote oldest visible summary (09:18) to the Old section to keep MAX 3 entries.
old_para = find_prefix(DEMOTE_PREFIX)
old_head = None
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("Old summary entries") or t == "## Old summary entries (tick history)":
        old_head = p
        break
if old_para is not None and old_head is not None:
    # prefix text with "OLD: "
    if old_para.runs:
        old_para.runs[0].text = "OLD: " + old_para.runs[0].text
    else:
        old_para.text = "OLD: " + old_para.text
    # physically move the paragraph to right after the Old-section heading (newest-at-top within OLD)
    old_head._p.addnext(old_para._p)
    demoted = True
else:
    demoted = False

# 3. Ensure one empty [] placeholder remains at the end of Open.
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
inprog = None
for p in d.paragraphs:
    if p.text.strip() == "In progress" and p.style.name.startswith("Heading"):
        inprog = p
        break
if not has_empty and inprog is not None:
    inprog.insert_paragraph_before("[]", style="normal")

d.save(DOC)
print("OK: 10:05 summary added; demoted 09:18:", demoted, "; empty [] present:", has_empty)
