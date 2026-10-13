#!/usr/bin/env python3
"""Backup ticker 2026-06-28 ~10:23 — update 'nothing happened' summary bullet."""
import docx
import sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

NEW_TEXT = (
    "2026-06-28 ~10:23 CDT (backup ticker): No new tasks or admin entries. "
    "Server healthy (HTTP 200). Library folder not synced to this machine — "
    "library doc not checked this tick. "
    "Both in-progress items (Alternative Ancients, remote access) still awaiting input."
)

CANDIDATE_PREFIXES = [
    "2026-06-28 ~09:52 CDT (backup ticker): No new tasks",
    "2026-06-28 ~09:52 CDT (backup ticker)",
    "2026-06-28 ~07:12 CDT (master ticker): No new tasks",
    "2026-06-28 ~09:22 CDT (master ticker)",
]

d = docx.Document(DOC)
target = None
for prefix in CANDIDATE_PREFIXES:
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            target = p
            break
    if target is not None:
        break

if target is None:
    print("ERROR: could not find 'nothing happened' bullet to update", file=sys.stderr)
    sys.exit(1)

if target.runs:
    target.runs[0].text = NEW_TEXT
    for extra in target.runs[1:]:
        extra.text = ""
else:
    target.add_run(NEW_TEXT)

d.save(DOC)
print(f"OK: updated master 'nothing happened' bullet to ~10:23 CDT")
