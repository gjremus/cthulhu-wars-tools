#!/usr/bin/env python3
"""Insert 5 open tasks into master current tasks doc."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

TASKS = [
    "[] Fix Ubbo-Sathla lingering pain indicator — unit shows 3 red pain lines but has not been in battle. Pain flag not being cleared. Game 540 (Horror vs Space). Investigate, fix, propagate to all builds, deploy.",
    "[] Fix Daemon Sultan showing 2 empty '?' spellbook slots — spellbooks were earned/satisfied but never picked. Related to the SBR/CG timing fix. Game 540 (Horror vs Space). Investigate, fix, propagate.",
    "[] Audit all recent updates (last 1-2 weeks) for rollback guide documentation — every code fix must have actual code details in the appropriate rollback guide.",
    "[] Fix MNU build git state — cw-mnu-wt has dirty/uncommitted state after CheckSpellbooksAction fix. Clean it up, commit, push.",
    "[]",
]

d = docx.Document(DOC)

# Find the "Open" section header
open_idx = None
for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text == "Open":
        open_idx = i
        break

if open_idx is None:
    print("ERROR: Could not find 'Open' section header")
    exit(1)

print(f"Found 'Open' at paragraph index {open_idx}")

# Find existing empty [] bullets after Open and before In Progress
in_progress_idx = None
empty_bullet_indices = []
for i in range(open_idx + 1, len(d.paragraphs)):
    text = d.paragraphs[i].text.strip()
    if text in ("In Progress", "In progress"):
        in_progress_idx = i
        break
    if text == "[]" or text == "[ ]":
        empty_bullet_indices.append(i)

print(f"Found {len(empty_bullet_indices)} empty bullets between Open and In Progress")
print(f"In Progress at index {in_progress_idx}")

# Strategy: replace existing empty bullets with our tasks, and add extras if needed
# First, replace existing empty bullets
replaced = 0
for idx, task_text in zip(empty_bullet_indices, TASKS):
    p = d.paragraphs[idx]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = task_text
    else:
        p.add_run(task_text)
    replaced += 1
    print(f"  Replaced bullet at idx {idx}: {task_text[:60]}...")

# If we have more tasks than empty bullets, insert new paragraphs
if replaced < len(TASKS):
    # Insert before In Progress
    from docx.oxml.ns import qn
    from copy import deepcopy

    # Get the last empty bullet we replaced as template for style
    template_idx = empty_bullet_indices[-1] if empty_bullet_indices else open_idx + 1
    template_p = d.paragraphs[template_idx]

    # Insert remaining tasks before in_progress paragraph
    ip_element = d.paragraphs[in_progress_idx]._element
    for task_text in TASKS[replaced:]:
        new_p = deepcopy(template_p._element)
        # Clear runs and set text
        for r in new_p.findall(qn('w:r')):
            new_p.remove(r)
        from docx.oxml import OxmlElement
        run_elem = OxmlElement('w:r')
        text_elem = OxmlElement('w:t')
        text_elem.text = task_text
        run_elem.append(text_elem)
        new_p.append(run_elem)
        ip_element.addprevious(new_p)
        print(f"  Inserted new bullet: {task_text[:60]}...")

d.save(DOC)
print("SAVED — all tasks inserted")
