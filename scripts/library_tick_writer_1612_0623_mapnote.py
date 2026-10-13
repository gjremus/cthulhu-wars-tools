#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

NOTE = (
    "  -- GOOD NEWS 2026-06-23 (investigation only, nothing deployed): a focused helper traced this to the "
    "exact spot and found the correction for the sideways view is ALREADY built into the Library edition's "
    "saved source (made on 2026-06-16) and double-checked correct against the real map artwork — units now "
    "sit in their proper rooms and the SL faction no longer vanishes from the Chamber of Sn'gac. Because "
    "the Library edition was rebuilt and republished on 2026-06-19, this correction has very likely already "
    "reached the live Library edition; I will confirm that the instant the site is reachable again (it is "
    "unreachable right now — see the site-down note above). What's still outstanding is the edition you were "
    "actually looking at (the More-Neutral-Units one): the same correction needs carrying across to it, "
    "which I am not allowed to do on my own, so this still needs your one-word OK to apply it there. There "
    "is also one tiny cosmetic leftover — in the sideways view a handful of rooms space their figures out a "
    "little less tidily than upright — but everything still lands in the correct room; I can polish that "
    "later if you'd like. Grovelling thanks for your patience."
)

d = docx.Document(DOC)

target = None
for p in d.paragraphs:
    t = p.text.strip()
    if "Mnu build" in t and "library 55 map" in t:
        target = p
        break
assert target is not None, "sideways-map task paragraph not found"

# Append the note as a new run that copies the formatting of the existing run.
if target.runs:
    base = target.runs[-1]
    r = target.add_run(NOTE)
    try:
        r.font.name = base.font.name
        r.font.size = base.font.size
        r.bold = base.bold
        r.italic = base.italic
    except Exception:
        pass
else:
    target.add_run(NOTE)

d.save(DOC)
print("OK: sideways-map task updated with 2026-06-23 good-news note.")
