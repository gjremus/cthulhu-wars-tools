#!/usr/bin/env python3
"""Add empty [] placeholder in Open section."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# Find the "In Progress" heading and insert [] just before it
in_progress_para = None
last_open_para = None
in_open = False

for p in d.paragraphs:
    text = p.text.strip()
    if text == "Open":
        in_open = True
        continue
    if text == "In Progress":
        in_progress_para = p
        break
    if in_open and text:
        last_open_para = p

# Check if empty [] already exists
has_empty = False
in_open2 = False
for p in d.paragraphs:
    text = p.text.strip()
    if text == "Open":
        in_open2 = True
        continue
    if text == "In Progress":
        break
    if in_open2 and text == "[]":
        has_empty = True

if not has_empty and in_progress_para is not None:
    # Insert [] just before "In Progress"
    new_p = OxmlElement('w:p')
    # Copy formatting from last open para if available
    ref = last_open_para if last_open_para else in_progress_para
    if ref._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(ref._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = "[]"
    run_elem.append(text_elem)
    new_p.append(run_elem)
    in_progress_para._element.addprevious(new_p)
    d.save(DOC)
    print("OK: empty [] added to Open section")
else:
    if has_empty:
        print("SKIP: empty [] already exists")
    else:
        print("ERROR: could not find In Progress heading")
