#!/usr/bin/env python3
"""Move the misplaced 'Cathedrals as GOOs (Holy Ground)' completed task
from the Faceless Blight Current Tasks doc to the Alt Ancients Current Tasks doc.

Word-for-word transfer. No summarization, no reproduction — exact text.
"""
import sys
import os
import shutil
import time
from docx import Document

CW = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars"
FBE_DOC = os.path.join(CW, "Factions/HomeBrews/Faceless Blight Current Tasks (1).docx")
ALT_DOC = os.path.join(CW, "Factions/HomeBrews/Alt Ancients Current Tasks.docx")
BACKUP_DIR = os.path.expanduser("~/.Trash")

MATCH_TEXT = "Cathedrals as GOOs (Holy Ground)"

def check_lockout(path):
    if not os.path.exists(path):
        return
    age = time.time() - os.path.getmtime(path)
    if age < 300:
        print(f"LOCKOUT: {os.path.basename(path)} modified {int(age)}s ago (need 300+). Aborting.", file=sys.stderr)
        sys.exit(1)

def backup(path):
    stamp = time.strftime("%Y%m%d-%H%M%S")
    dst = os.path.join(BACKUP_DIR, f"{os.path.basename(path)}.pre_move_{stamp}")
    shutil.copy2(path, dst)
    print(f"Backup: {dst}")

def find_and_remove_from_fbe():
    """Find the single misplaced paragraph in FBE doc, remove it, return its text."""
    doc = Document(FBE_DOC)

    found_idx = None
    found_text = None

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if MATCH_TEXT in text:
            found_idx = i
            found_text = para.text
            print(f"FOUND in FBE at paragraph {i}:")
            print(f"  {text[:150]}...")
            break

    if found_idx is None:
        print("ERROR: Could not find the misplaced entry in FBE doc!", file=sys.stderr)
        sys.exit(1)

    # Remove the single paragraph
    para = doc.paragraphs[found_idx]
    p_element = para._element
    p_element.getparent().remove(p_element)
    print(f"  Removed paragraph {found_idx} from FBE doc")

    doc.save(FBE_DOC)
    print(f"FBE doc saved (entry removed).")
    return [found_text]

def append_to_alt_ancients(texts):
    """Append the entry word-for-word to the Alt Ancients Completed section."""
    doc = Document(ALT_DOC)

    # Find the "Completed" section
    completed_idx = None
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip().lower() == "completed":
            completed_idx = i
            break

    if completed_idx is None:
        print("ERROR: Could not find 'Completed' section in Alt Ancients doc!", file=sys.stderr)
        sys.exit(1)

    print(f"Found 'Completed' section at paragraph {completed_idx} in Alt Ancients doc")

    # Find insertion point: right after "Completed" header, before existing completed items
    # Insert after the first bullet placeholder "- []" if it exists, or right after the header
    insert_after = completed_idx
    for i in range(completed_idx + 1, min(completed_idx + 5, len(doc.paragraphs))):
        text = doc.paragraphs[i].text.strip()
        if text == "- []" or text == "[]":
            insert_after = i
            break
        elif text.startswith("- ") or text.startswith("☒"):
            # There's already content, insert before it
            insert_after = i - 1
            break

    # Insert the text right after the placeholder or header
    # We need to add a new paragraph after insert_after
    from docx.oxml.ns import qn
    from copy import deepcopy

    # Insert after insert_after position
    ref_para = doc.paragraphs[insert_after]._element
    parent = ref_para.getparent()

    # Create new paragraphs for each text line
    for text_line in texts:
        new_para = doc.add_paragraph()
        new_para.text = text_line
        # Move it to the right position (after ref_para)
        new_element = new_para._element
        parent.remove(new_element)
        ref_para.addnext(new_element)
        ref_para = new_element  # Next insertion goes after this one
        print(f"  Inserted into Alt Ancients: {text_line[:80]}...")

    doc.save(ALT_DOC)
    print(f"Alt Ancients doc saved (entry added to Completed).")

def main():
    print("=" * 60)
    print("MOVE TASK: Alt Ancients entry from FBE doc → Alt Ancients doc")
    print("=" * 60)
    print(f"Looking for: '{MATCH_TEXT}'")
    print()

    # Safety checks
    check_lockout(FBE_DOC)
    check_lockout(ALT_DOC)

    # Backups
    backup(FBE_DOC)
    backup(ALT_DOC)

    # Step 1: Find and remove from FBE
    print("\n--- Step 1: Remove from FBE doc ---")
    texts = find_and_remove_from_fbe()

    # Step 2: Append to Alt Ancients
    print("\n--- Step 2: Append to Alt Ancients doc ---")
    append_to_alt_ancients(texts)

    # Step 3: Verify
    print("\n--- Step 3: Verify ---")
    verify_doc = Document(FBE_DOC)
    for para in verify_doc.paragraphs:
        if MATCH_TEXT in para.text:
            print("WARNING: Entry still found in FBE doc after removal!", file=sys.stderr)
            sys.exit(1)
    print("  FBE doc: entry confirmed removed.")

    verify_doc2 = Document(ALT_DOC)
    found = False
    for para in verify_doc2.paragraphs:
        if MATCH_TEXT in para.text:
            found = True
            break
    if not found:
        print("WARNING: Entry not found in Alt Ancients doc after insertion!", file=sys.stderr)
        sys.exit(1)
    print("  Alt Ancients doc: entry confirmed present.")

    print("\nDONE. Task moved word-for-word.")

if __name__ == '__main__':
    main()
