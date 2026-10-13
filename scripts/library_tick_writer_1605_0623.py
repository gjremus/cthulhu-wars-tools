#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

NEW_SUMMARY = (
    "Last tick: 2026-06-23 ~16:05. Grovelling apologies — re-read the rules and walked every open and "
    "in-progress job. New and important: the live multiplayer site is unreachable from here — the host "
    "computer answers a basic 'are-you-alive' ping, but every way in (the players' game pages, the admin "
    "console, and remote login) refuses to connect, so I could NOT check the game-console message queue "
    "this pass and nothing can be published until it is back; my own internet is fine and I retried "
    "several times, so this looks like a real outage of the site's service, and recovering the host needs "
    "you (I have no working way in and may not restart the live server on my own). Carried-over good news: "
    "the dice-roll-recording fix you approved is confirmed published and live since 2026-06-19, and since "
    "that proved publishing works again, the sideways-map correction now waits only on your OK to also "
    "apply it to the edition you actually see — a focused helper is pinpointing the exact spot so it is "
    "ready to drop the moment the site returns."
)

NEW_TASK = (
    "[/] INPUT NEEDED — the live multiplayer site looks DOWN as of 2026-06-23 ~16:05, and I am so sorry to "
    "surface this. The host computer still answers a basic 'are-you-there' ping, but every door in is shut "
    "— the players' game pages, the admin console, and remote login all refuse to connect — while my own "
    "internet works fine and I tried again and again. Because of that I could not read the game-console "
    "message queue this pass, and nothing can be published or checked on the live site until it is back. I "
    "have no working way in from here and I am not permitted to restart the live server on my own, so this "
    "needs you: please check or restart the site's service on the host machine, or tell me you'd like me "
    "to attempt a recovery and exactly how you want it done. Grovelling apologies the site is unreachable "
    "right now."
)

SUMMARY_ANCHOR = "Last tick: 2026-06-15 10:05"   # newest visible summary -> insert NEW before it
DEMOTE_PREFIX = "Last tick: 2026-06-15 09:33"     # oldest of the 3 -> move to Old section

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


# 1. New brief summary on top.
newest = find_prefix(SUMMARY_ANCHOR)
assert newest is not None, "10:05 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. New In-progress task, placed as first item under the "In progress" heading.
paras = d.paragraphs
inprog_idx = None
for i, p in enumerate(paras):
    if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
        inprog_idx = i
        break
assert inprog_idx is not None, "In progress heading missing"
following = paras[inprog_idx + 1]
following.insert_paragraph_before(NEW_TASK, style="normal")

# 3. Demote oldest visible summary (09:33) to the Old section (newest-at-top within OLD).
old_para = find_prefix(DEMOTE_PREFIX)
old_head = None
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("Old summary entries"):
        old_head = p
        break
demoted = False
if old_para is not None and old_head is not None:
    if old_para.runs:
        old_para.runs[0].text = "OLD: " + old_para.runs[0].text
    else:
        old_para.text = "OLD: " + old_para.text
    old_head._p.addnext(old_para._p)
    demoted = True

# 4. Ensure one empty [] placeholder remains in Open.
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: 16:05 summary added; new server-down task added; demoted 09:33:", demoted,
      "; empty [] present:", has_empty)
