#!/usr/bin/env python3
"""Backup ticker 2026-06-29 ~15:23 — update existing nothing-happened bullet timestamp."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

OLD_PREFIX = "2026-06-29 ~15:10 CDT (backup ticker): No changes this tick"
NEW_TEXT = ("2026-06-29 ~15:23 CDT (backup ticker): No changes this tick — "
            "server reachable, admin queue empty (0 entries). Both in-progress "
            "tasks remain INPUT NEEDED (Alt Ancients awaiting user answers, "
            "Telegram bot awaiting token). Library doc inaccessible "
            "(Google Drive resource deadlock on entire folder).")

found = False
for p in d.paragraphs:
    if p.text.strip().startswith(OLD_PREFIX):
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = NEW_TEXT
        else:
            p.add_run(NEW_TEXT)
        found = True
        break

if not found:
    print("ERROR: could not find the existing nothing-happened bullet to update")
    raise SystemExit(1)

d.save(DOC)
print("OK: nothing-happened bullet timestamp updated to ~15:23")
