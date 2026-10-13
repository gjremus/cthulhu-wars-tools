#!/usr/bin/env python3
"""
library_tick_writer_0803_0625.py
Library ticker pass, 2026-06-25 ~08:05.
1. Create .bak copy next to docx (oldest already trashed).
2. Insert new summary bullet (08:05) before the current newest (06:03).
3. Demote the two oldest summary bullets (~05:34, ~05:49) to Old summary entries
   so the summary has exactly 2 nothing-actionable bullets (per Rules 20 + 21).
4. Regenerate the .last_seen_tasks.md mirror.
"""
import sys
import shutil
import time
import os
from docx import Document

DOC = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
    "library build current tasks.docx"
)
MIRROR = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/.last_seen_tasks.md"
)

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)
print(f"Backup: {os.path.basename(bak)}")

d = Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~08:05. Grovelling apologies — re-read the full rules, drained "
    "the admin message box (empty, nothing new), and confirmed the live site fully healthy "
    "(the Library single-player edition at the main address, the More-Neutral-Units beta, "
    "and the admin console all answered instantly). Walked every open and waiting job; "
    "nothing was changed or deployed, and everything still waits only on your go-aheads "
    "(chiefly the sideways-map correction for the edition you actually see and the crashed "
    "five-player game fix)."
)

# Anchors for the two oldest summary bullets to demote
ANCHOR_0534 = "Last tick: 2026-06-25 ~05:34"
ANCHOR_0549 = "Last tick: 2026-06-25 ~05:49"
# Anchor for the bullet we insert before (keep this one, it becomes bullet #2)
ANCHOR_0603 = "Last tick: 2026-06-25 ~06:03"

# Find all needed paragraphs
p_0534 = None
p_0549 = None
p_0603 = None
old_heading = None

for p in d.paragraphs:
    txt = p.text.strip()
    if txt.startswith(ANCHOR_0534):
        p_0534 = p
    elif txt.startswith(ANCHOR_0549):
        p_0549 = p
    elif txt.startswith(ANCHOR_0603):
        p_0603 = p
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p

if p_0603 is None:
    print("ERROR: 06:03 summary anchor not found", file=sys.stderr)
    sys.exit(1)
if old_heading is None:
    print("ERROR: Old summary entries heading not found", file=sys.stderr)
    sys.exit(1)

# Step 1: Insert the new 08:05 summary bullet before the 06:03 bullet
p_0603.insert_paragraph_before(NEW_SUMMARY, style=p_0603.style)
print("Inserted new summary (08:05) before 06:03 bullet")

# Step 2: Demote 05:49 to OLD (if found)
if p_0549 is not None:
    if not p_0549.text.strip().startswith("OLD:"):
        set_text(p_0549, "OLD: " + p_0549.text.strip())
    old_heading._p.addnext(p_0549._p)
    print("Demoted 05:49 bullet to Old section")
else:
    print("WARNING: 05:49 anchor not found — may already have been demoted")

# Step 3: Demote 05:34 to OLD (if found)
if p_0534 is not None:
    if not p_0534.text.strip().startswith("OLD:"):
        set_text(p_0534, "OLD: " + p_0534.text.strip())
    old_heading._p.addnext(p_0534._p)
    print("Demoted 05:34 bullet to Old section")
else:
    print("WARNING: 05:34 anchor not found — may already have been demoted")

# Step 4: Save
d.save(DOC)
print(f"OK: saved {DOC}")

# Step 5: Regenerate .last_seen_tasks.md mirror
lines = []
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lines.append("## " + p.text)
    else:
        lines.append(p.text)
with open(MIRROR, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Mirror regenerated: {len(lines)} paragraphs")
