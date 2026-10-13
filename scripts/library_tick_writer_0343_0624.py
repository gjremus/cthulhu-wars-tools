#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

# Brief new tick summary (<=3 sentences), grovelling, layman's terms.
NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~03:43. Grovelling apologies — re-read the full rules, drained the game "
    "console's message box (empty, nothing new), and confirmed the live site healthy (the Library "
    "single-player edition at the main address, the More-Neutral-Units beta and the admin console all "
    "answered). I walked every open and waiting job; nothing was deployed or changed live, and "
    "everything substantial still waits only on your go-aheads."
)

ANCHOR_0149 = "Last tick: 2026-06-24 ~01:49"   # newest visible summary
DEMOTE_0146 = "Last tick: 2026-06-24 ~01:46"   # oldest visible nothing-actionable -> OLD section

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


# 1. Insert brief new summary on top (before the 01:49 line).
newest = find_prefix(ANCHOR_0149)
assert newest is not None, "01:49 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. Demote the 01:46 line to the Old section (newest-at-top within OLD) to keep
#    nothing-actionable summary bullets at the allowed maximum of two.
old_para = find_prefix(DEMOTE_0146)
old_head = None
for p in d.paragraphs:
    if p.text.strip().startswith("Old summary entries"):
        old_head = p
        break
demoted = False
if old_para is not None and old_head is not None:
    if old_para.runs:
        old_para.runs[0].text = "OLD: " + old_para.runs[0].text
    else:
        old_para.text = "OLD: " + old_para.text
    old_head._p.addnext(old_para._p)
    demoted = True

# 3. Ensure one empty [] placeholder remains in Open (right before the In progress heading).
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: 03:43 summary added; demoted 01:46:", demoted,
      "; empty [] present:", has_empty)
