#!/usr/bin/env python3
"""
library_tick_writer_0848_0625.py
Library ticker pass, 2026-06-25 ~08:48 CDT.
1. Create .bak copy next to docx.
2. Insert new summary bullet at top of Bulleted summary.
3. Demote the "Last tick: 2026-06-25 ~08:18" nothing-actionable bullet to Old summary entries
   (per Rule 21: max 2 nothing-actionable bullets; oldest must move out when adding a new one).
4. Regenerate the .last_seen_tasks.md mirror.
"""
import sys
import shutil
import time
import os
from docx import Document

DOC = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
    "library build current tasks.docx"
)
MIRROR = (
    "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
    "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/.last_seen_tasks.md"
)

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)
print(f"Backup: {os.path.basename(bak)}")

d = Document(DOC)


def set_text(p, txt):
    runs = list(p.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(txt)


NEW_SUMMARY = (
    "Last tick: 2026-06-25 ~08:48. Grovelling apologies — re-read the full rules, drained "
    "the admin message box (empty, nothing new), and confirmed the live site fully healthy "
    "(the Library single-player edition, the More-Neutral-Units beta and the admin console "
    "all answered instantly). I walked every open and waiting job; nothing was changed or "
    "deployed, and everything still waits only on your go-aheads (chiefly the sideways-map "
    "correction for the edition you actually see and the crashed five-player game fix)."
)

# Anchor: the most recent existing summary bullet (08:37) — insert BEFORE it
ANCHOR_0837 = "Last tick: 2026-06-25 ~08:37"

# Anchor: the nothing-actionable bullet to demote (08:18)
ANCHOR_0818 = "Last tick: 2026-06-25 ~08:18"

# ── Step 1: Find the 08:37 paragraph ─────────────────────────────────────────
anchor_0837_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR_0837):
        anchor_0837_p = p
        break

if anchor_0837_p is None:
    print("ERROR: 08:37 summary anchor not found", file=sys.stderr)
    sys.exit(1)

# ── Step 2: Insert new summary before the 08:37 paragraph ────────────────────
anchor_0837_p.insert_paragraph_before(NEW_SUMMARY, style=anchor_0837_p.style)
print("Inserted new summary before 08:37 bullet")

# ── Step 3: Find the 08:18 paragraph and demote to Old summary entries ────────
anchor_0818_p = None
for p in d.paragraphs:
    if p.text.strip().startswith(ANCHOR_0818):
        anchor_0818_p = p
        break

if anchor_0818_p is None:
    print("ERROR: 08:18 summary anchor not found", file=sys.stderr)
    sys.exit(1)

old_heading = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_heading = p
        break

if old_heading is None:
    print("ERROR: Old summary entries heading not found", file=sys.stderr)
    sys.exit(1)

if not anchor_0818_p.text.strip().startswith("OLD:"):
    set_text(anchor_0818_p, "OLD: " + anchor_0818_p.text.strip())
old_heading._p.addnext(anchor_0818_p._p)
print("Demoted 08:18 bullet to Old section")

# ── Step 4: Save ──────────────────────────────────────────────────────────────
d.save(DOC)
print(f"OK: saved {DOC}")

# ── Step 5: Regenerate .last_seen_tasks.md mirror ────────────────────────────
lines = []
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lines.append("## " + p.text)
    else:
        lines.append(p.text)
with open(MIRROR, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Mirror regenerated: {len(lines)} paragraphs")
