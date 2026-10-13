#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~19:12 — update idle summary bullet timestamp."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

OLD_FRAGMENT = "~19:00 CDT"
NEW_TEXT = "2026-06-29 ~19:12 CDT: Both in-progress tasks still blocked on user input (Alt Ancients awaiting Section 1 answers, Telegram awaiting bot token). Admin queue empty."

found = False
for p in d.paragraphs:
    if OLD_FRAGMENT in p.text and "blocked on user input" in p.text:
        for run in p.runs:
            run.text = ""
        p.runs[0].text = NEW_TEXT
        found = True
        break

if not found:
    for p in d.paragraphs:
        if "blocked on user input" in p.text and "2026-06-29" in p.text:
            for run in p.runs:
                run.text = ""
            p.runs[0].text = NEW_TEXT
            found = True
            break

if not found:
    print("WARNING: could not find the target bullet to update")
else:
    d.save(DOC)
    print("OK: idle summary bullet timestamp updated to ~19:12")
