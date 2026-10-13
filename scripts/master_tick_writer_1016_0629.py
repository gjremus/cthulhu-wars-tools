#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~10:16 CDT — Library build investigation complete."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# === 1. Update the Library build task ===
# Find the paragraph containing "library build" in Open section
for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if "library build" in text.lower() and "what the fuck" in text.lower():
        # Clear existing runs and replace with new text
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = (
                "[/] INPUT NEEDED — Library build: I have verified the live server. "
                "The correct Library code IS deployed (confirmed by file hash match, "
                "no Bubastis references anywhere in the client JS or server JAR, "
                "HTML title says library-at-celaeno-v4). The game log for "
                "cixsasvawrtztikq is clean — all Library factions, no parse errors, "
                "no Bubastis actions. The wrong-code issue was already fixed earlier today. "
                "If you are still seeing a crash, please try a hard refresh "
                "(Ctrl+Shift+R / Cmd+Shift+R) to clear browser cache. "
                "If it still crashes after that, what error do you see? "
                "(a blank screen, a specific error message, the game freezing at a certain point?)"
            )
        break

# === 2. Move the Library task from Open to In Progress ===
# Find the paragraph we just modified, and move it after "In Progress" heading
# Actually, the task now has [/] so it should be in the In Progress section.
# Find "In Progress" heading and the library task
lib_para = None
lib_para_elem = None
in_progress_heading = None

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text == "In Progress":
        in_progress_heading = p
    if "INPUT NEEDED" in text and "Library build" in text:
        lib_para = p
        lib_para_elem = p._element

# Move library task to after "In Progress" heading (if it's not already there)
if lib_para_elem is not None and in_progress_heading is not None:
    # Check if it's currently BEFORE the In Progress heading (i.e., in Open section)
    body = d.element.body
    lib_idx = list(body).index(lib_para_elem)
    ip_idx = list(body).index(in_progress_heading._element)
    if lib_idx < ip_idx:
        # Remove from current position
        body.remove(lib_para_elem)
        # Insert after In Progress heading
        in_progress_heading._element.addnext(lib_para_elem)

# === 3. Update summary bullets ===
# Find "Bulleted summary" heading and manage the 3-bullet limit
in_summary = False
summary_bullets = []

for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text and not in_summary:
        in_summary = True
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
        elif text in ("Open", "In Progress", "Completed") or text.startswith("["):
            break

# Find "Old summary bullets" section
old_summary_heading = None
for p in d.paragraphs:
    if "Old summary bullets" in p.text:
        old_summary_heading = p
        break

# If we have 3+ summary bullets, move the oldest to old summaries
if len(summary_bullets) >= 3 and old_summary_heading is not None:
    oldest_para = summary_bullets[-1][1]  # last = oldest (newest at top)
    oldest_text = oldest_para.text.strip()

    # Create new paragraph for old summaries
    new_p = OxmlElement('w:p')
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_heading._element.addnext(new_p)

    # Remove the oldest from summary section
    oldest_para._element.getparent().remove(oldest_para._element)

# Add new summary bullet at the top
# Re-find first summary bullet after modifications
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

new_summary_text = (
    "2026-06-29 ~10:16 CDT (master ticker): Library build VERIFIED CORRECT — "
    "no Bubastis code anywhere on the live server (confirmed by file hash, "
    "content search, HTML title). The wrong-code issue was already fixed earlier today. "
    "Marked INPUT NEEDED: asking if crash persists after hard refresh. "
    "MNU tasks and Prime Cause fix still blocked on compilation (no Java on this machine). "
    "Admin queue unreachable (server endpoint timing out)."
)

if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = new_summary_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)

d.save(DOC)
print("OK: Library task updated to INPUT NEEDED, summary added")
