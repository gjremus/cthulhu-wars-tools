#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~12:55 — update summary: move oldest to old-summaries, add nothing-happened tick."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# Find the 3 summary bullets
in_summary = False
summary_bullets = []
for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text:
        in_summary = True
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
        elif text in ("Open", "In Progress", "Completed", "[]", "[/]", "[x]"):
            break
        elif text.startswith("["):
            break

# Find old summary bullets section
old_summary_para = None
for p in d.paragraphs:
    if "Old summary bullets" in p.text:
        old_summary_para = p
        break

if len(summary_bullets) >= 3 and old_summary_para is not None:
    # Move the oldest (last) bullet to old summaries section
    oldest_para = summary_bullets[-1][1]
    oldest_text = oldest_para.text.strip()

    # Insert after old summary heading
    new_p = OxmlElement('w:p')
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_para._element.addnext(new_p)

    # Remove oldest from summary
    oldest_para._element.getparent().remove(oldest_para._element)

# Now add the new "nothing happened" bullet at the top of summary
# Re-find first summary bullet
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

NEW_BULLET = ("2026-06-29 ~12:55 CDT (master ticker): No changes this tick — "
              "server unreachable (DNS timeout), admin queue inaccessible. "
              "Both in-progress tasks remain INPUT NEEDED (Alt Ancients awaiting "
              "user answers, Telegram bot awaiting server).")

if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = NEW_BULLET
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)

d.save(DOC)
print("OK: summary updated — oldest moved to old-summaries, nothing-happened tick added at top")
