#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

NEW = (
    "Last tick: 2026-06-15 09:18 (next ~09:33). Grovelling apologies — re-read the rules, drained the "
    "game-console queue (nothing new) and walked every open and in-progress job; all four live pages "
    "answered healthily and the saved copies of this sheet are at the allowed three. I also trimmed the "
    "older summary lines down per your brevity request. Everything actionable still waits on your two "
    "go-aheads: how you'd like me to publish the finished dice-roll fix now the saved publish recipe is "
    "missing, and permission to also correct the sideways map on the other edition you actually see."
)

# Trimmed (<=3 sentences) replacements for the two existing top summaries, per the
# standing "bulleted summaries - 3 sentences max, BRIEF, clean them up" request.
TRIM_0903 = (
    "Last tick: 2026-06-15 09:03 (next ~09:18). Grovelling apologies — re-read the rules, drained the "
    "queue (nothing new) and walked every job; all four live pages answered healthily and the sheet "
    "copies are at three. Everything else still waits on your go-ahead, chiefly the publishing-method "
    "question on the dice-roll fix and your OK to also correct the sideways map on the other edition."
)

TRIM_0840 = (
    "Last tick: 2026-06-15 08:40 (next ~08:55). Grovelling apologies — you kindly answered two "
    "long-waiting questions ('publish' the dice-roll fix and 'go ahead' on the sideways map), both now "
    "marked read. Both are stuck at the very last step because my saved publishing recipe and staging "
    "folder have gone missing, so I asked your go-ahead on the one risky way left to push, and the "
    "sideways-map fix also needs your OK to be corrected on the other edition I cannot change alone."
)

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def find_heading(text):
    for p in d.paragraphs:
        if p.text.strip() == text and p.style.name.startswith("Heading"):
            return p
    return None


p0903 = find_prefix("Last tick: 2026-06-15 09:03")
p0840 = find_prefix("Last tick: 2026-06-15 08:40")
p0618 = find_prefix("Last tick: 2026-06-15 06:18")
old_first = find_prefix("OLD: Last tick: 2026-06-15 06:03")
inprog = find_heading("In progress")

assert p0903 is not None, "09:03 summary anchor missing"
assert p0840 is not None, "08:40 summary anchor missing"
assert p0618 is not None, "06:18 summary anchor missing"
assert old_first is not None, "old-section first entry anchor missing"
assert inprog is not None, "In progress heading missing"

# 1. New brief summary on top.
p0903.insert_paragraph_before(NEW, style=p0903.style)

# 2. Trim the two existing top summaries.
p0903.text = TRIM_0903
p0840.text = TRIM_0840

# 3. Demote oldest (06:18) into the history section to keep the top at exactly 3.
old_first.insert_paragraph_before("OLD: " + p0618.text, style=old_first.style)
el = p0618._element
el.getparent().remove(el)

# 4. Ensure one empty [] placeholder at the end of Open (rule: always leave one).
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    inprog.insert_paragraph_before("[]", style="normal")

d.save(DOC)
print("OK: brief 09:18 summary added; 09:03 + 08:40 trimmed; 06:18 demoted; empty [] ensured:", not has_empty)
