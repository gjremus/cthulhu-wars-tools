#!/usr/bin/env python3
import docx, shutil, time, os

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

# In-dir pre-write backup (matches existing convention).
bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~05:18. Grovelling apologies — the live multiplayer site has gone DOWN AGAIN, "
    "the second seizure in under half an hour: the server machine still answers a basic network ping, but "
    "its website and remote-login services both freeze partway through every connection, the classic sign "
    "it has run out of memory or processing room and locked up. I cannot get in to restart it (remote login "
    "is frozen the same way) and the message queue could not be read this pass, so it again needs a hard "
    "reboot from your cloud control panel — the full ask is in the in-progress list below. I changed and "
    "deployed nothing."
)

NEW_INPROG = (
    "[/] INPUT NEEDED — the live multiplayer site is DOWN AGAIN as of 2026-06-24 ~05:18, and this is now the "
    "SECOND time it has seized today (it briefly recovered around 05:03, then locked up again). Grovelling "
    "apologies. The pattern is exactly the same as before: the server computer still answers a basic network "
    "ping, and it even accepts the very first knock on both the website door and the remote-login door, but "
    "nothing behind those doors ever finishes replying — both the website software and the remote-login "
    "service freeze midway through every connection. That signature almost always means the machine has run "
    "out of memory, processing power, or disk space and seized up; because the recovery did not last, "
    "something is steadily refilling whatever keeps filling. I cannot get in to restart anything (remote "
    "login is frozen the same way) and I could not read the admin message box this pass either. INPUT "
    "NEEDED: please reboot the server from your cloud hosting control panel again when you can; the moment it "
    "is back I will confirm every page loads and then dig into what keeps exhausting it so this stops "
    "recurring. I changed nothing and deployed nothing this pass."
)

ANCHOR_0503 = "Last tick: 2026-06-24 ~05:03"          # current newest summary bullet
ANCHOR_0343 = "Last tick: 2026-06-24 ~03:43"          # oldest displayed status bullet -> demote
INPROG_HEAD = "In progress"

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


# 1. New brief summary on top (before the 05:03 line); inherit that bullet's (normal) style.
newest = find_prefix(ANCHOR_0503)
assert newest is not None, "05:03 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. Demote oldest displayed status bullet (03:43) into Old summary entries to keep MAX 3.
old_head = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Old summary entries" in p.text:
        old_head = p
        break
demoted = False
older = find_prefix(ANCHOR_0343)
if older is not None and old_head is not None:
    if older.runs:
        older.runs[0].text = "OLD: " + older.runs[0].text
    else:
        older.text = "OLD: " + older.text
    old_head._p.addnext(older._p)
    demoted = True

# 3. New INPUT-NEEDED in-progress item at top of In progress section.
inprog_head = None
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and p.text.strip() == INPROG_HEAD:
        inprog_head = p
        break
assert inprog_head is not None, "In progress heading missing"
# find first paragraph after the heading to inherit its style and insert before it
first_item = None
seen = False
for p in d.paragraphs:
    if p._p is inprog_head._p:
        seen = True
        continue
    if seen and p.text.strip():
        first_item = p
        break
if first_item is not None:
    first_item.insert_paragraph_before(NEW_INPROG, style=first_item.style)
else:
    inprog_head.insert_paragraph_before(NEW_INPROG, style="normal")

# 4. Ensure one empty [] placeholder remains in Open.
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == INPROG_HEAD:
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: summary added; 03:43 demoted:", demoted,
      "; new in-progress added; empty [] present:", has_empty,
      "; backup:", os.path.basename(bak))
