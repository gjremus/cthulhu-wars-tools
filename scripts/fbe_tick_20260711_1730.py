#!/usr/bin/env python3
"""FBE tick update 2026-07-11 ~17:30:
Consolidate two idle tick lines into one (per Rule 30).
Remove 16:15 line, update 17:15 line to 17:30.
Admin queue drained, no FBE-specific issues. Server healthy."""
import zipfile
import shutil
import tempfile
import os
import docx
from docx.oxml.ns import qn
import docx.oxml

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

def patch_content_types(docx_path):
    """Add Default Extension='ttf' to [Content_Types].xml if not present."""
    tmp = tempfile.mktemp(suffix='.docx')
    shutil.copy2(docx_path, tmp)

    with zipfile.ZipFile(tmp, 'r') as zin:
        ct = zin.read('[Content_Types].xml').decode('utf-8')
        if 'Extension="ttf"' not in ct:
            ct = ct.replace(
                '<Default Extension="odttf"',
                '<Default Extension="ttf" ContentType="application/x-font-ttf" /><Default Extension="odttf"'
            )
            tmp2 = tempfile.mktemp(suffix='.docx')
            with zipfile.ZipFile(tmp2, 'w', zipfile.ZIP_DEFLATED) as zout:
                for item in zin.namelist():
                    if item == '[Content_Types].xml':
                        zout.writestr(item, ct.encode('utf-8'))
                    else:
                        zout.writestr(item, zin.read(item))
            os.remove(tmp)
            return tmp2
    return tmp

patched = patch_content_types(PATH)
d = docx.Document(patched)

# Step 1: Remove the 16:15 idle tick line
removed_16 = False
for i, p in enumerate(d.paragraphs):
    if "2026-07-11 ~16:15" in p.text and "no open tasks" in p.text.lower():
        parent = p._element.getparent()
        parent.remove(p._element)
        print(f"Removed 16:15 idle tick line (paragraph {i})")
        removed_16 = True
        break

# Step 2: Update the 17:15 idle tick line to 17:30
updated_17 = False
for i, p in enumerate(d.paragraphs):
    if "2026-07-11 ~17:15" in p.text and "no open tasks" in p.text.lower():
        new_text = "[2026-07-11 ~17:30] Most humbly, deepest apologies: admin queue drained (no new FBE-specific issues). Server healthy, no open tasks remain. Your worthless servant grovels abjectly."
        for child in list(p._element):
            if child.tag.endswith('}r'):
                p._element.remove(child)
        run_elem = docx.oxml.OxmlElement('w:r')
        text_elem = docx.oxml.OxmlElement('w:t')
        text_elem.text = new_text
        text_elem.set(qn('xml:space'), 'preserve')
        run_elem.append(text_elem)
        p._element.append(run_elem)
        print(f"Updated 17:15 idle tick to 17:30 (paragraph {i})")
        updated_17 = True
        break

if not removed_16 and not updated_17:
    print("WARNING: No matching tick lines found — no write performed.")
    os.remove(patched)
    raise SystemExit(1)

if not removed_16:
    print("NOTE: 16:15 line not found (may have been removed already)")

if not updated_17:
    print("WARNING: 17:15 line not found — could not update timestamp")
    os.remove(patched)
    raise SystemExit(1)

d.save(PATH)
os.remove(patched)
print("FBE task doc saved.")
