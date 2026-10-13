#!/usr/bin/env python3
import re, sys
import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor

DOC = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/HomeBrews/The Burrowers Beneath (TB) Faction Implementation Guide.docx"
CONTENT = "/Users/gremus/Claude-Projects/cthulhu-wars-tools/scripts/tb_section2_content.md"
QUESTIONS = "/Users/gremus/Claude-Projects/cthulhu-wars-tools/scripts/tb_section2_questions.txt"
RED = RGBColor(0xC0, 0x00, 0x00)

# ---- parse content by '### <token> ...' ----
sections = {}  # token -> list of body lines
cur = None
for line in open(CONTENT, encoding="utf-8"):
    line = line.rstrip("\n")
    m = re.match(r"^###\s+(\S+)\s", line)
    if m:
        cur = m.group(1)
        sections[cur] = []
    elif cur is not None and line.strip():
        sections[cur].append(line)

d = docx.Document(DOC)
paras = d.paragraphs

def is_heading(p):
    return p.style.name in ("Heading 1", "Heading 2")

def first_token(t):
    t = t.strip()
    return t.split()[0] if t else ""

# locate Section 2 / Section 3 boundaries
s2 = s3 = None
for i, p in enumerate(paras):
    t = p.text.strip()
    if t.startswith("SECTION 2"):
        s2 = i
    elif t.startswith("SECTION 3"):
        s3 = i
        break
assert s2 is not None and s3 is not None, "section bounds not found"

def new_para_after(anchor_p, text):
    np_el = OxmlElement("w:p")
    anchor_p.addnext(np_el)
    np = Paragraph(np_el, None)
    if text:
        np.add_run(text)
    return np

filled = []
i = s2 + 1
while i < s3:
    p = paras[i]
    if p.style.name == "Heading 2":
        tok = first_token(p.text)
        if tok in sections:
            # collect old body (until next heading)
            j = i + 1
            old_body = []
            while j < len(paras) and not is_heading(paras[j]):
                old_body.append(paras[j])
                j += 1
            # insert new content after heading, in order
            cursor = p._p
            for line in sections[tok]:
                npg = new_para_after(cursor, line)
                cursor = npg._p
            # delete old stub body
            for ob in old_body:
                ob._p.getparent().remove(ob._p)
            filled.append(tok)
    i += 1

# ---- open questions block: insert before SECTION 3 heading ----
# re-find SECTION 3 element (indices shifted after edits)
s3_p = None
for p in d.paragraphs:
    if p.text.strip().startswith("SECTION 3"):
        s3_p = p
        break
assert s3_p is not None

qlines = [l.rstrip("\n") for l in open(QUESTIONS, encoding="utf-8") if l.strip()]
for idx, line in enumerate(qlines):
    qp_el = OxmlElement("w:p")
    s3_p._p.addprevious(qp_el)
    qp = Paragraph(qp_el, None)
    run = qp.add_run(line)
    run.font.color.rgb = RED
    if idx == 0:
        run.font.bold = True

# ---- update status A.iii ----
status_done = False
for p in d.paragraphs:
    if p.text.strip() == "Section 2 Status: Empty":
        for r in list(p.runs):
            r.text = ""
        if p.runs:
            p.runs[0].text = "Section 2 Status: Claude Audited"
        else:
            p.add_run("Section 2 Status: Claude Audited")
        status_done = True
        break

# ---- A.vii pointer note ----
note = ("[2026-06-11 guide ticker] Section 2 generated and self-audited against Section 1 + the "
        "General Faction Creation Guide and existing-faction conventions (BB Moon tile, multi-part GOO, "
        "RoA tintFilter/overlay/glyph standards). 12 open questions are posted in red at the end of "
        "Section 2 for creator response. Section 2 status set to Claude Audited, awaiting Creator Responded.")
for p in d.paragraphs:
    if p.text.strip().startswith("Clean out questions after they are answered"):
        new_para_after(p._p, note)
        break

d.save(DOC)
print("FILLED subsections:", filled)
print("status updated:", status_done)
print("questions inserted:", len(qlines))
