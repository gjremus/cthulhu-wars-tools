#!/usr/bin/env python3
"""Append the TB Mantle HUD cultist centering fix entry to the Homebrew Rollback Guide."""
import docx

PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Homebrew Deployed Online Rollback Guide.docx"

d = docx.Document(PATH)
before = len(d.paragraphs)

def h1(t): d.add_paragraph(t, style="Heading 1")
def h2(t): d.add_paragraph(t, style="Heading 2")
def body(t): d.add_paragraph(t, style="Body Text")

h1("Fix HB-TB-Mantle: Mantle HUD popup no longer centers a Cultist image when the gate is abandoned (2026-07-05)")

h2("Symptom.")
body("When the gate on the Mantle was abandoned (no cultist controlling it), the Mantle popup could still show a Cultist image centered on the gate if any Cadavolyte or Acolyte happened to be present. The creator requested no centered cultist when the gate is uncontrolled.")

h2("Root cause.")
body("The overlay code unconditionally treated any Cadavolyte/Acolyte on the Mantle as a gate controller whenever a gate was present, without checking whether any faction actually controlled that gate.")

h2("Fix.")
body("CthulhuWarsSolo.scala now checks whether the Mantle gate is actively controlled (any faction has it in their gates list) and passes 'controlled' or 'abandoned' in the gate entry string. overlay.scala reads this marker and only centers a cultist as the gate controller when the gate is confirmed controlled. When abandoned, cultists are arranged in the outer ring with other units instead.")

h2("Files changed.")
body("solo/CthulhuWarsSolo.scala — added mantleGateControlled check; changed gate entry label from 'Gate (TB)' to 'Gate (controlled)' or 'Gate (abandoned)'. solo/overlay.scala — added gateControlled flag check; cultists only go to centered gateControllerEntries when gate is controlled; when abandoned, all units including cultists go to the ring layout.")

h2("Rollback.")
body("In solo/CthulhuWarsSolo.scala: revert the mantleGateControlled variable and restore the gateEntry to unconditionally use 'gate|Gate (TB)|76'. In solo/overlay.scala: revert gateControlled logic, restore gateControllerEntries to unconditionally filter Cadavolyte/Acolyte when gateEntries.nonEmpty, remove abandonedCultistEntries, and restore otherUnits to the original filter. Rebuild and redeploy.")

d.save(PATH)
after = len(d.paragraphs)
print("Appended TB Mantle cultist fix entry. Paragraphs:", before, "->", after, "(+%d)" % (after - before))
