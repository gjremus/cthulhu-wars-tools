#!/usr/bin/env python3
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Library at Celaeno/"
       "library build current tasks.docx")

# Brief new tick summary (<=3 sentences), grovelling, layman's terms.
NEW_SUMMARY = (
    "Last tick: 2026-06-24 ~01:49. Grovelling apologies — re-read the full rules, found the game "
    "console's message box empty, and confirmed the live site healthy (the Library single-player "
    "edition at the main address, the More-Neutral-Units beta and the admin console all answered). "
    "Nothing was deployed or changed live; everything substantial still waits only on your go-aheads, "
    "and the saved copies of this sheet stay at the allowed three."
)

# Brief rewrite of the prior visible summary, honouring the "keep summaries brief" request.
REWRITE_0146 = (
    "Last tick: 2026-06-24 ~01:46. Grovelling apologies — re-read the rules, found the console "
    "message box empty, and confirmed all live pages healthy. Nothing was deployed; the substantive "
    "jobs still wait only on your go-aheads."
)

ANCHOR_0146 = "Last tick: 2026-06-24 ~01:46"   # newest visible summary
DEMOTE_2134 = "Last tick: 2026-06-23 ~21:34"   # oldest visible nothing-actionable -> OLD section

d = docx.Document(DOC)


def find_prefix(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


# 1. Insert brief new summary on top (before the 01:46 line).
newest = find_prefix(ANCHOR_0146)
assert newest is not None, "01:46 summary anchor missing"
newest.insert_paragraph_before(NEW_SUMMARY, style=newest.style)

# 2. Trim the 01:46 line to its brief form.
set_text(newest, REWRITE_0146)

# 3. Demote the 21:34 line to the Old section (newest-at-top within OLD).
old_para = find_prefix(DEMOTE_2134)
old_head = None
for p in d.paragraphs:
    if p.text.strip().startswith("Old summary entries"):
        old_head = p
        break
demoted = False
if old_para is not None and old_head is not None:
    if old_para.runs:
        old_para.runs[0].text = "OLD: " + old_para.runs[0].text
    else:
        old_para.text = "OLD: " + old_para.text
    old_head._p.addnext(old_para._p)
    demoted = True

# 4. Ensure one empty [] placeholder remains in Open (right before the In progress heading).
has_empty = any(p.text.strip() == "[]" for p in d.paragraphs)
if not has_empty:
    for p in d.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == "In progress":
            p.insert_paragraph_before("[]", style="normal")
            break

d.save(DOC)
print("OK: 01:49 summary added; 01:46 trimmed; demoted 21:34:", demoted,
      "; empty [] present:", has_empty)
