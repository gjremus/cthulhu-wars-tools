#!/usr/bin/env python3
# Backup ticker 2026-06-26 ~13:08 — safety-net pass on library doc.
# Per MAX 2 "NOTHING ACTIONABLE" rule: update existing backup-ticker bullet in place.
import docx
import shutil
import time
import sys

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

NEW_SUMMARY = (
    "Last tick: 2026-06-26 ~13:08 (backup ticker). Grovelling apologies "
    "— re-read the full rules, drained the admin message box (empty, nothing new), "
    "and confirmed the live site healthy (Library and More-Neutral-Units both "
    "answered instantly). Walked every open and waiting job; nothing was changed "
    "or deployed, and everything still waits only on your go-aheads (chiefly the "
    "sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

# Find the existing backup ticker summary bullet to update in place
CANDIDATE_PREFIXES = [
    "Last tick: 2026-06-26 ~12:52 (backup ticker)",
    "Last tick: 2026-06-26 ~11:07 (backup ticker)",
    "Last tick: 2026-06-26 ~10:50 (backup ticker)",
]

target = None
for prefix in CANDIDATE_PREFIXES:
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            target = p
            break
    if target is not None:
        break

if target is None:
    print("ERROR: could not find backup-ticker summary bullet to update", file=sys.stderr)
    sys.exit(1)

if target.runs:
    target.runs[0].text = NEW_SUMMARY
    for extra in target.runs[1:]:
        extra.text = ""
else:
    target.add_run(NEW_SUMMARY)

d.save(DOC)
print("WRITE COMPLETE")
