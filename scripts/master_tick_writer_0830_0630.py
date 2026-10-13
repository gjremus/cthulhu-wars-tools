#!/usr/bin/env python3
"""Master ticker 2026-06-30 ~08:30 — move MN to in-progress INPUT NEEDED, update summary."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# === TASK 1: Move []MN from Open to In Progress with INPUT NEEDED ===
open_idx = None
in_progress_idx = None
mn_para = None

for i, p in enumerate(d.paragraphs):
    text = p.text.strip()
    if text == "Open":
        open_idx = i
    elif text == "In Progress":
        in_progress_idx = i
    if open_idx and not in_progress_idx:
        if text == "[]MN" or text == "[] MN" or text == "[]MN " or text.replace(" ", "") == "[]MN":
            mn_para = p
            break

if mn_para:
    # Rewrite the paragraph with INPUT NEEDED
    for run in mn_para.runs:
        run.text = ""
    mn_para.runs[0].text = ('[/] INPUT NEEDED: "MN" task — what does this refer to? '
                            'Please clarify what you would like done so I can start work on it.')

    # Move from Open to In Progress: remove from current, insert after In Progress heading
    elem = mn_para._element
    parent = elem.getparent()
    parent.remove(elem)

    # Find In Progress heading and insert after it (before existing in-progress items)
    for p in d.paragraphs:
        if p.text.strip() == "In Progress":
            p._element.addnext(elem)
            break
    print("OK: MN moved to In Progress with INPUT NEEDED")
else:
    print("NOTE: MN paragraph not found in Open section (may have been removed/changed)")

# === TASK 2: Ensure there is still an empty [] in Open ===
# Re-scan after move
open_found = False
in_progress_found = False
has_empty_bullet = False
last_open_para = None

for p in d.paragraphs:
    text = p.text.strip()
    if text == "Open":
        open_found = True
        continue
    if text == "In Progress":
        in_progress_found = True
    if open_found and not in_progress_found:
        last_open_para = p
        if text == "[]" or text == "[ ]":
            has_empty_bullet = True

if not has_empty_bullet and last_open_para is not None:
    # Insert an empty [] bullet after the last Open para (before In Progress)
    new_p = OxmlElement('w:p')
    if last_open_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(last_open_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = "[]"
    run_elem.append(text_elem)
    new_p.append(run_elem)
    last_open_para._element.addnext(new_p)
    print("OK: added empty [] to Open section")
elif has_empty_bullet:
    print("OK: empty [] already exists in Open")

# === TASK 3: Update bulleted summary (3-bullet max, newest at top) ===
in_summary = False
summary_bullets = []

for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text and not in_summary:
        in_summary = True
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
        elif text in ("Open", "In Progress", "Completed") or text.startswith("["):
            break

# Find "Old summary bullets" section
old_summary_para = None
for p in d.paragraphs:
    if "Old summary bullets" in p.text or "old summary bullets" in p.text.lower():
        old_summary_para = p
        break

# If we already have 3 summary bullets, move the oldest to old-summaries
if len(summary_bullets) >= 3 and old_summary_para is not None:
    oldest_para = summary_bullets[-1][1]
    oldest_text = oldest_para.text.strip()

    # Create new paragraph in old summaries
    new_p = OxmlElement('w:p')
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_para._element.addnext(new_p)

    # Remove from summary section
    oldest_para._element.getparent().remove(oldest_para._element)
    print(f"OK: moved oldest summary bullet to old-summaries")

# Re-scan to find insert point for new summary bullet
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

NEW_SUMMARY = ("2026-06-30 ~08:30 CDT (master ticker): Both in-progress tasks still "
               "blocked on user input (Alt Ancients awaiting Section 1 answers, Telegram "
               "awaiting bot token). Moved ambiguous \"MN\" task to in-progress with INPUT "
               "NEEDED for clarification. Admin queue empty. All 5 builds healthy (200).")

if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = NEW_SUMMARY
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)
    print("OK: new summary bullet inserted")
else:
    print("WARNING: could not find summary insert point")

d.save(DOC)
print("SAVED")
