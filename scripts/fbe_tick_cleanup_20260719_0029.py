#!/usr/bin/env python3
"""Clean up duplicate 'nothing happened' tick lines per rule 30."""
import sys
from docx import Document

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

# The pattern to look for (nothing-happened ticks)
NOTHING_PATTERN = "admin queue inaccessible (token file locked). No open or in-progress"

# The one tick we want to keep (most recent)
KEEP_TICK = "[2026-07-19 ~00:28] Most humbly, deepest apologies: admin queue inaccessible (token file locked). No open or in-progress FBE tasks remain. Your worthless servant grovels abjectly."

doc = Document(DOC_PATH)
found_anchor = False
found_keep = False
to_remove = []

for i, para in enumerate(doc.paragraphs):
    if "Bulleted summary (newest at top)" in para.text:
        found_anchor = True
        continue

    if found_anchor:
        # Stop when we hit the next section
        if para.text.strip() and para.text.strip()[0].isupper() and "~" not in para.text:
            # This might be the next section heading
            if len(para.text) > 50 and "[" not in para.text:
                break

        # Check if this is a nothing-happened tick
        if NOTHING_PATTERN in para.text:
            if KEEP_TICK in para.text:
                found_keep = True
            else:
                # Mark for removal
                to_remove.append(i)

# Remove the duplicate paragraphs (in reverse order to preserve indices)
for idx in reversed(to_remove):
    p = doc.paragraphs[idx]._element
    p.getparent().remove(p)

print(f"Removed {len(to_remove)} duplicate 'nothing happened' tick lines")
doc.save(DOC_PATH)
