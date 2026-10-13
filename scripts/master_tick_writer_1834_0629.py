#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~18:34 — nothing-happened tick summary update."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

# 1. Find the "Bulleted summary" heading and the summary bullets after it
in_summary = False
summary_bullets = []
summary_heading_idx = None

for i, p in enumerate(d.paragraphs):
    if "Bulleted summary" in p.text and not in_summary:
        in_summary = True
        summary_heading_idx = i
        continue
    if in_summary:
        text = p.text.strip()
        if text.startswith("202"):
            summary_bullets.append((i, p))
        elif text in ("Open", "In Progress", "Completed") or text == "":
            if text in ("Open", "In Progress", "Completed"):
                break
        elif text.startswith("["):
            break

# 2. Move the oldest (3rd) summary bullet to old summaries section
# Find "Old summary bullets" section
old_summary_idx = None
for i, p in enumerate(d.paragraphs):
    if "Old summary bullets" in p.text:
        old_summary_idx = i
        break

if len(summary_bullets) >= 3 and old_summary_idx is not None:
    oldest_para = summary_bullets[2][1]  # 3rd bullet (0-indexed)
    oldest_text = oldest_para.text.strip()

    # Find the first paragraph after "Old summary bullets" heading
    old_summary_para = d.paragraphs[old_summary_idx]
    # Insert the oldest bullet right after the old-summary heading
    # by inserting a new element after the heading
    new_p = OxmlElement('w:p')
    # Copy paragraph properties from oldest_para for formatting
    if oldest_para._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(oldest_para._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    # Add text run
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = oldest_text
    run_elem.append(text_elem)
    new_p.append(run_elem)
    old_summary_para._element.addnext(new_p)

    # Remove the oldest from the summary section
    oldest_para._element.getparent().remove(oldest_para._element)

# 3. Add "nothing happened" bullet at the top of the summary section
# Insert before the first remaining summary bullet
# Re-scan after the removal
first_summary = None
in_summary2 = False
for p in d.paragraphs:
    if "Bulleted summary" in p.text:
        in_summary2 = True
        continue
    if in_summary2 and p.text.strip().startswith("202"):
        first_summary = p
        break

if first_summary is not None:
    new_p = OxmlElement('w:p')
    if first_summary._element.find(qn('w:pPr')) is not None:
        new_pPr = deepcopy(first_summary._element.find(qn('w:pPr')))
        new_p.insert(0, new_pPr)
    run_elem = OxmlElement('w:r')
    text_elem = OxmlElement('w:t')
    text_elem.text = "2026-06-29 ~18:34 CDT (master ticker): No new work this tick. Both in-progress tasks still waiting on user input (Alternative Ancients clarifications, remote access preference). Admin queue empty."
    run_elem.append(text_elem)
    new_p.append(run_elem)
    first_summary._element.addprevious(new_p)

d.save(DOC)
print("OK: nothing-happened summary added, oldest summary moved to old-summaries")
