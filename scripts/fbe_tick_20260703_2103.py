#!/usr/bin/env python3
"""
FBE tick writer — 2026-07-03 ~21:03
Task: Nothing new. Admin queue empty, no user responses. Build verified clean.
In-progress task still awaiting deploy permission. Insert one idle tick line.
"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

doc = Document(DOC_PATH)
paras = doc.paragraphs

# Find 'Bulleted summary' section
summary_header_idx = None
for i, p in enumerate(paras):
    if "Bulleted summary" in p.text:
        summary_header_idx = i
        break

if summary_header_idx is None:
    print("ERROR: could not find 'Bulleted summary' paragraph")
    sys.exit(1)

print(f"Found 'Bulleted summary' at index {summary_header_idx}")

# Find first bullet and count
first_bullet_idx = None
bullet_count = 0
for i in range(summary_header_idx + 1, len(paras)):
    text = paras[i].text.strip()
    if text.startswith('[2026-'):
        if first_bullet_idx is None:
            first_bullet_idx = i
        bullet_count += 1
    elif text and not text.startswith('[2026-'):
        break

print(f"Found {bullet_count} existing summary bullets, first at index {first_bullet_idx}")

# Show current bullets
for i in range(summary_header_idx, min(summary_header_idx + 8, len(paras))):
    print(f"  [{i}]: {paras[i].text[:120]}")

NEW_TICK_LINE = "[2026-07-03 ~21:03] Most humbly: tick verified build compiles clean with the Animated Rush legacy fix. Still awaiting your gracious permission to deploy (games 530, 534). Admin queue empty."

def clone_para_with_text(source_para, new_text):
    """Clone a paragraph's formatting and replace its text."""
    new_para = deepcopy(source_para._element)
    for r in new_para.findall(qn('w:r')):
        new_para.remove(r)
    r = OxmlElement('w:r')
    source_runs = source_para._element.findall(qn('w:r'))
    if source_runs:
        rpr = source_runs[0].find(qn('w:rPr'))
        if rpr is not None:
            r.append(deepcopy(rpr))
    t = OxmlElement('w:t')
    t.text = new_text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    new_para.append(r)
    return new_para

# Insert new tick bullet BEFORE the first existing bullet
new_tick_elem = clone_para_with_text(paras[first_bullet_idx], NEW_TICK_LINE)
paras[first_bullet_idx]._element.addprevious(new_tick_elem)

# If we now have more than 6 summary bullets, remove the oldest
paras_after = doc.paragraphs
bullet_indices = []
for i in range(summary_header_idx + 1, len(paras_after)):
    text = paras_after[i].text.strip()
    if text.startswith('[2026-'):
        bullet_indices.append(i)
    elif text and not text.startswith('[2026-'):
        break

if len(bullet_indices) > 6:
    for idx in bullet_indices[6:]:
        paras_after[idx]._element.getparent().remove(paras_after[idx]._element)
    print(f"Trimmed {len(bullet_indices) - 6} oldest summary bullet(s)")

# Verify final state
paras_final = doc.paragraphs
print(f"\nFinal summary section:")
for i in range(summary_header_idx, min(summary_header_idx + 10, len(paras_final))):
    print(f"  [{i}]: {paras_final[i].text[:140]}")

doc.save(DOC_PATH)
print("\nDOC SAVED.")
