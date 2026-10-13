#!/usr/bin/env python3
"""
FBE tick writer — 2026-07-03 ~21:05
Move the [/] INPUT NEEDED task from Open section to In Progress section.
(Rule 10: Move [/] in progress tasks to the In Progress section immediately.)
"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy

DOC_PATH = "/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/My Drive/Personal/Games/Cthulhu Wars/Factions/HomeBrews/Faceless Blight Current Tasks (1).docx"

doc = Document(DOC_PATH)
paras = doc.paragraphs

# Find section headers
open_header_idx = None
in_progress_header_idx = None
completed_header_idx = None

for i, p in enumerate(paras):
    text = p.text.strip()
    if text == "Open" and open_header_idx is None:
        open_header_idx = i
    if text == "In Progress" and in_progress_header_idx is None:
        in_progress_header_idx = i
    if text == "Completed" and completed_header_idx is None:
        completed_header_idx = i

print(f"Open at: {open_header_idx}")
print(f"In Progress at: {in_progress_header_idx}")
print(f"Completed at: {completed_header_idx}")

if not all([open_header_idx, in_progress_header_idx]):
    print("ERROR: Could not find required sections")
    sys.exit(1)

# Find the [/] INPUT NEEDED task in Open section
task_idx = None
for i in range(open_header_idx + 1, in_progress_header_idx):
    if "[/] INPUT NEEDED" in paras[i].text:
        task_idx = i
        break

if task_idx is None:
    print("Task not found in Open section — may already be moved. Exiting cleanly.")
    sys.exit(0)

print(f"Found [/] INPUT NEEDED task at index {task_idx}")

# Clone the paragraph element
task_element = deepcopy(paras[task_idx]._element)

# Remove from Open section
paras[task_idx]._element.getparent().remove(paras[task_idx]._element)
print("Removed from Open section")

# Re-find In Progress header (indices shifted after removal)
paras = doc.paragraphs
in_progress_header_idx = None
completed_header_idx = None
for i, p in enumerate(paras):
    text = p.text.strip()
    if text == "In Progress" and in_progress_header_idx is None:
        in_progress_header_idx = i
    if text == "Completed" and completed_header_idx is None:
        completed_header_idx = i

print(f"In Progress now at: {in_progress_header_idx}")

# Insert after In Progress header (before Completed)
paras[in_progress_header_idx]._element.addnext(task_element)
print("Inserted into In Progress section")

doc.save(DOC_PATH)
print("\nDOC SAVED.")
