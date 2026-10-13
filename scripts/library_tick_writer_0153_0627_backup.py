#!/usr/bin/env python3
# Library backup ticker 2026-06-27 ~01:53 — routine safety-net pass, nothing changed.
# Update the existing summary bullet from ~00:51 with fresh timestamp.
import docx
import shutil
import time
import os
import glob

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
       "library build current tasks.docx")

bak = DOC + ".bak.pre_write_%d" % int(time.time())
shutil.copy2(DOC, bak)

d = docx.Document(DOC)

OLD_ANCHOR = "Last tick: 2026-06-27 ~00:51 (backup ticker)"
NEW_SUMMARY = (
    "Last tick: 2026-06-27 ~01:53 (backup ticker). Grovelling apologies — "
    "re-read the full rules, drained the admin message box (empty, nothing new), "
    "and confirmed the live site healthy (Library and More-Neutral-Units both "
    "answered instantly). Walked every open and waiting job; nothing was changed "
    "or deployed, and everything still waits only on your go-aheads (chiefly the "
    "sideways-map correction for the More-Neutral-Units edition you actually see, "
    "and the crashed five-player game fix)."
)

found = False
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith(OLD_ANCHOR):
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = NEW_SUMMARY
        else:
            p.add_run(NEW_SUMMARY)
        found = True
        print("Updated existing backup ticker summary bullet timestamp")
        break

if not found:
    print("WARNING: Could not find existing backup ticker bullet to update")

# Rotate backups: keep max 3 in the main dir
bak_pattern = DOC + ".bak.pre_write_*"
all_baks = sorted(glob.glob(bak_pattern))
trash_dir = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
             "My Drive/Personal/Games/Cthulhu Wars/Maps/Library at Celaeno/"
             ".Trash-bak-rotation/")
os.makedirs(trash_dir, exist_ok=True)
while len(all_baks) > 3:
    oldest = all_baks.pop(0)
    dest = os.path.join(trash_dir, os.path.basename(oldest))
    if os.path.exists(dest):
        dest = dest + ".%d" % int(time.time())
    os.rename(oldest, dest)
    print("Rotated to Trash-bak-rotation: %s" % os.path.basename(oldest))

d.save(DOC)
print("OK: library doc tick updated")
