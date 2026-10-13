#!/usr/bin/env python3
"""Master ticker 2026-06-29 ~18:28 — update summary bullet: admin queue now reachable and empty."""
import docx

DOC = ("/Users/gremus/Library/CloudStorage/GoogleDrive-gremus@salesforce.com/"
       "My Drive/Personal/Games/Cthulhu Wars/master current tasks.docx")

d = docx.Document(DOC)

OLD_TEXT = "2026-06-29 ~18:11 CDT: Both in-progress tasks blocked on user input (Alt Ancients awaiting Section 1 answers, Telegram awaiting bot token). Admin queue unreachable (server up but endpoint returning 404)."
NEW_TEXT = "2026-06-29 ~18:28 CDT: Both in-progress tasks still blocked on user input (Alt Ancients awaiting Section 1 answers, Telegram awaiting bot token). Admin queue empty."

found = False
for p in d.paragraphs:
    if "2026-06-29 ~18:11 CDT" in p.text and "blocked on user input" in p.text:
        full = p.text
        for run in p.runs:
            run.text = ""
        p.runs[0].text = full.replace(OLD_TEXT, NEW_TEXT)
        found = True
        break

if not found:
    for p in d.paragraphs:
        if "~18:11 CDT" in p.text:
            for run in p.runs:
                run.text = ""
            p.runs[0].text = NEW_TEXT
            found = True
            break

if not found:
    print("WARNING: could not find the target bullet to update")
else:
    d.save(DOC)
    print("OK: summary bullet updated — admin queue now healthy")
